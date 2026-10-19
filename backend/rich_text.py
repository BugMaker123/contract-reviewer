"""Convert pasted, untrusted HTML to a local DOCX without loading external content."""
import re
from lxml import html
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .word_export import paragraphs, visible_paragraph_text

SKIP = {'script', 'style', 'iframe', 'object', 'embed', 'svg', 'math', 'img', 'head', 'link', 'meta'}
BLOCKS = {'p', 'div', 'section', 'article', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'blockquote', 'pre', 'ul', 'ol', 'table', 'li'}


def styles(node, inherited=None):
    result = dict(inherited or {})
    result.update(dict(part.split(':', 1) for part in (node.get('style') or '').split(';') if ':' in part))
    result = {key.strip().lower(): value.strip() for key, value in result.items()}
    tag = str(node.tag).lower()
    if tag in ('b', 'strong'): result['font-weight'] = 'bold'
    if tag in ('i', 'em'): result['font-style'] = 'italic'
    if tag == 'u': result['text-decoration'] = 'underline'
    if tag == 'th': result['font-weight'] = 'bold'
    if tag == 'font':
        sizes = {'1':'8pt','2':'10pt','3':'12pt','4':'14pt','5':'18pt','6':'24pt','7':'36pt'}
        if node.get('size') in sizes: result['font-size'] = sizes[node.get('size')]
        if node.get('color'): result['color'] = node.get('color')
        if node.get('face'): result['font-family'] = node.get('face')
    return result


def color(value):
    value = value.strip().lower()
    names = {'red':'ff0000','blue':'0000ff','black':'000000','green':'008000','gray':'808080','white':'ffffff'}
    value = names.get(value, value.lstrip('#'))
    if re.fullmatch('[0-9a-f]{3}', value): value = ''.join(c*2 for c in value)
    if re.fullmatch('[0-9a-f]{6}', value): return RGBColor.from_string(value.upper())
    match = re.fullmatch(r'rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)', value)
    if match and all(0 <= int(v) <= 255 for v in match.groups()): return RGBColor(*map(int, match.groups()))


def add_text(paragraph, text, style):
    if not text: return
    run = paragraph.add_run(text)
    if style.get('font-weight', '').lower() in ('bold','bolder','600','700','800','900'): run.bold = True
    if style.get('font-style', '').lower() == 'italic': run.italic = True
    if 'underline' in style.get('text-decoration', ''): run.underline = True
    match = re.match(r'^(\d+(?:\.\d+)?)(pt|px)?$', style.get('font-size', ''))
    if match:
        size = float(match[1]) * (.75 if match[2] == 'px' else 1)
        run.font.size = Pt(max(6, min(72, size)))
    if style.get('font-family'):
        name = style['font-family'].split(',')[0].strip(' "' + "'")[:80]
        run.font.name = name
        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    if style.get('color'):
        value = color(style['color'])
        if value: run.font.color.rgb = value
    if style.get('background-color'):
        value = color(style['background-color'])
        if value:
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), str(value)); run._r.get_or_add_rPr().append(shading)


def inline(node, paragraph, inherited=None):
    if str(node.tag).lower() in SKIP: return
    style = styles(node, inherited)
    add_text(paragraph, node.text, style)
    for child in node:
        if not isinstance(child.tag, str): continue
        tag = str(child.tag).lower()
        if tag == 'br': add_text(paragraph, '\n', style)
        elif tag not in ('ul','ol','table'):
            inline(child, paragraph, style)
        add_text(paragraph, child.tail, style)


def format_paragraph(paragraph, node, inherited=None):
    style = styles(node, inherited)
    align = style.get('text-align', node.get('align', '')).lower()
    if align in ('left','right','center','justify'):
        paragraph.alignment = {'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT,'center':WD_ALIGN_PARAGRAPH.CENTER,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    if str(node.tag).lower() == 'blockquote': paragraph.paragraph_format.left_indent = Inches(.3)
    inline(node, paragraph, style)


def table(node, parent):
    rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
    if not rows: return
    placements, occupied, columns = [], set(), 0
    if len(rows) > 200: raise ValueError('粘贴表格超过200行，请改用上传DOCX')
    for y, row in enumerate(rows):
        x = 0
        for cell in row:
            if str(cell.tag).lower() not in ('td','th'): continue
            while (y,x) in occupied: x += 1
            try: width, height = int(cell.get('colspan', 1)), int(cell.get('rowspan', 1))
            except ValueError: raise ValueError('表格合并单元格格式无效')
            if not (1 <= width <= 30 and 1 <= height <= len(rows)-y and x+width <= 30): raise ValueError('表格结构过于复杂，请改用上传DOCX')
            placements.append((y,x,height,width,cell))
            for yy in range(y,y+height):
                for xx in range(x,x+width): occupied.add((yy,xx))
            columns = max(columns,x+width); x += width
    if not columns: return
    output = parent.add_table(rows=len(rows), cols=columns); output.style = 'Table Grid'
    for y,x,height,width,cell_node in placements:
        cell = output.cell(y,x)
        if height > 1 or width > 1: cell = cell.merge(output.cell(y+height-1,x+width-1))
        initial = cell.paragraphs[0]
        render_children(cell_node, cell)
        if len(cell.paragraphs) > 1 and not initial.text: initial._p.getparent().remove(initial._p)


def render_children(node, parent, inherited=None):
    style = styles(node, inherited)
    current = None
    if node.text and node.text.strip():
        current = parent.add_paragraph(); add_text(current,node.text,style)
    for child in node:
        if not isinstance(child.tag, str): continue
        tag = str(child.tag).lower()
        if tag in SKIP: pass
        elif tag == 'table': table(child,parent); current = None
        elif tag in ('ul','ol'):
            for item in child:
                if str(item.tag).lower() != 'li': continue
                paragraph = parent.add_paragraph(style='List Bullet' if tag == 'ul' else 'List Number')
                format_paragraph(paragraph,item,style)
                for nested in item:
                    if str(nested.tag).lower() in ('ul','ol'):
                        wrapper = html.Element('div'); wrapper.append(nested); render_children(wrapper,parent,style)
            current = None
        elif tag in BLOCKS:
            if any(str(c.tag).lower() in BLOCKS for c in child) and tag not in ('p','li','pre'):
                render_children(child,parent,style)
            else:
                paragraph = parent.add_paragraph(style=('Heading '+tag[1]) if tag in ('h1','h2','h3','h4','h5','h6') else None)
                format_paragraph(paragraph,child,style)
            current = None
        else:
            if current is None: current = parent.add_paragraph()
            inline(child,current,style)
        if child.tail and child.tail.strip():
            if current is None: current = parent.add_paragraph()
            add_text(current,child.tail,style)


def pasted_document(rich_html, plain_text=''):
    document = Document()
    if rich_html.strip():
        root = html.fragment_fromstring(rich_html, create_parent='div', parser=html.HTMLParser(no_network=True))
        if sum(1 for _ in root.iter()) > 20000: raise ValueError('粘贴内容过于复杂，请改用上传DOCX')
        render_children(root,document)
    else:
        for line in plain_text.split('\n'): document.add_paragraph(line)
    text = '\n'.join(visible_paragraph_text(p._p) for p in paragraphs(document))
    if not text.strip(): raise ValueError('请粘贴可读取的合同正文')
    if len(text) > 200000: raise ValueError('合同超过20万字，请拆分后审查')
    return document, text
