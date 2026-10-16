import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm
from zipfile import ZipFile
from fastapi.testclient import TestClient

from backend.main import app
from backend import main, workspace_api, graph
from backend.store import Repository
from backend.sample_contracts import SAMPLE_CONTRACTS


@pytest.fixture
def client(tmp_path, monkeypatch):
    repository = Repository(tmp_path / 'workspace.db')
    monkeypatch.setattr(workspace_api, 'repo', repository)
    monkeypatch.setattr(main, 'repository', repository)
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: None)
    with TestClient(app) as instance:
        yield instance


def create_task(client, **changes):
    body = {'title': '内部测试合同', 'type': '采购合同', 'text': '第一条 标的\n交付硬件产品\n第二条 付款\n验收后30日付款。', 'mode': 'quick', **changes}
    response = client.post('/api/tasks', json=body)
    assert response.status_code == 200, response.text
    return response.json()


def review(client, task):
    response = client.post('/api/tasks/' + task['id'] + '/review/stream', json={})
    assert response.status_code == 200, response.text
    blocks = response.text.strip().split('\n\n')
    assert len([b for b in blocks if b.startswith('event: node_update')]) == 5
    completed = next(block for block in blocks if block.startswith('event: complete'))
    return json.loads(completed.split('data: ', 1)[1])


def test_multiple_contracts_and_versions_survive_reopen(client):
    first = create_task(client, title='合同一')
    second = create_task(client, title='合同二')
    run1 = review(client, first)
    task = client.get('/api/tasks/' + first['id']).json()
    response = client.put('/api/tasks/' + first['id'], json={**task, 'text': task['text'] + '\n新增违约条款', 'revision': task['revision']})
    assert response.status_code == 200
    run2 = review(client, first)
    assert run1['run_id'] != run2['run_id']
    assert client.get('/api/runs/' + run1['run_id']).json()['input']['text'] == first['text']
    assert '新增违约条款' in client.get('/api/runs/' + run2['run_id']).json()['input']['text']
    # New repository connection reads the same persisted objects, without any browser session.
    repository = Repository(workspace_api.repo.path)
    assert len(repository.list('task')) == 2
    assert repository.get(second['id'])['title'] == '合同二'
    assert len(repository.list('run', first['id'])) == 2


def test_optimistic_conflict_and_recoverable_archive(client):
    task = create_task(client)
    assert client.put('/api/tasks/' + task['id'], json={**task, 'title': '新名称'}).status_code == 200
    assert client.put('/api/tasks/' + task['id'], json={**task, 'title': '过期页面的名称'}).status_code == 409
    latest = client.get('/api/tasks/' + task['id']).json()
    archived = client.post('/api/tasks/' + task['id'] + '/archive', json={'archived': True, 'revision': latest['revision']}).json()
    assert archived['archived']
    assert client.post('/api/tasks/' + task['id'] + '/review/stream', json={}).status_code == 409
    restored = client.post('/api/tasks/' + task['id'] + '/archive', json={'archived': False, 'revision': archived['revision']}).json()
    assert not restored['archived']


def checklist(client):
    return client.post('/api/checklists', json={'name': '采购企业要求', 'published': True, 'items': [
        {'id': 'payment', 'title': '付款期限', 'requirement': '必须约定30日付款', 'operator': 'contains_all', 'keywords': ['30日', '付款']},
        {'id': 'invoice', 'title': '发票', 'requirement': '必须约定专票', 'operator': 'contains_any', 'keywords': ['专用发票']},
        {'id': 'balance', 'title': '责任对等', 'requirement': '责任须对等', 'operator': 'semantic'}]}).json()


def test_checklist_snapshot_and_honest_pending_status(client):
    standard = checklist(client)
    task = create_task(client, mode='checklist', checklist_id=standard['id'])
    changed = client.put('/api/checklists/' + standard['id'], json={**standard, 'items': [standard['items'][0]], 'revision': standard['revision']}).json()
    assert changed['version'] == 2
    result = review(client, task)['data']
    assert [item['status'] for item in result['checklist_findings']] == ['pass', 'fail', 'needs_review']
    saved = client.get('/api/tasks/' + task['id']).json()
    assert saved['run']['input']['checklist']['version'] == 1
    assert len(saved['run']['input']['checklist']['items']) == 3
    assert result['summary']['compliance_score'] is None


def test_compare_only_reports_changed_clause(client):
    reference = '第一条 标的\n提供硬件产品\n第二条 付款\n验收后30日付款。\n第三条 保密\n双方保密。'
    task = create_task(client, mode='compare', text=reference.replace('30日', '60日'), reference_text=reference)
    result = review(client, task)['data']
    assert len(result['differences']) == 1
    assert result['differences'][0]['original'] == '验收后30日付款。'
    assert result['differences'][0]['current'] == '验收后60日付款。'
    assert result['summary']['compliance_score'] is None


def test_background_and_scale_reach_graph_prompt(client, monkeypatch):
    prompts = []
    class FakeModel:
        def invoke(self, messages):
            prompts.append(messages[-1].content)
            return type('Response', (), {'content': json.dumps({'tax_settlement': {}, 'issues': []})})()
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: FakeModel())
    task = create_task(client, mode='deep', background='我方为采购方，要求分期交付', scale='强势', focus='验收条件', references=[{'name': '企业政策', 'text': '付款需经财务确认'}])
    review(client, task)
    assert prompts and all(word in prompts[0] for word in ('强势', '分期交付', '验收条件', '财务确认'))


def test_comparison_highlights_only_changed_characters(client):
    original = '双方依照法律履行本协议。\n验收后30日付款，其余条款保持不变。'
    revised = '双方依照法律履行本协议。\n验收后60日付款，其余条款保持不变。'
    result = client.post('/api/documents/compare', json={'original':original,'revised':revised}).json()
    assert ''.join(original[a:b] for a,b in result['original_chars']) == '3'
    assert ''.join(revised[a:b] for a,b in result['revised_chars']) == '6'


def test_handling_survives_and_is_bound_to_the_correct_run(client):
    sample = SAMPLE_CONTRACTS['business_cooperation_contract']
    task = create_task(client, text=sample['text'])
    completed = review(client, task)
    run = client.get('/api/runs/' + completed['run_id']).json()
    first = run['result']['issues'][0]
    assert first['redline']['original_text'] in task['text']
    handling = {'revision': run['revision'], 'applied': [first['id']], 'comments': {first['id']: '业务已确认'}, 'decisions': {}}
    response = client.put('/api/runs/' + run['id'] + '/handling', json=handling)
    assert response.status_code == 200
    saved = response.json()
    assert saved['comments'][first['id']] == '业务已确认'
    assert saved['decisions'][first['id']] == 'applied'
    assert client.put('/api/runs/' + run['id'] + '/handling', json=handling).status_code == 409
    invalid = {**handling, 'revision': saved['revision'], 'applied': ['foreign_issue']}
    assert client.put('/api/runs/' + run['id'] + '/handling', json=invalid).status_code == 422
    document = client.get('/api/runs/' + run['id'] + '/export?scope=current')
    assert document.status_code == 200
    text = '\n'.join(p.text for p in Document(BytesIO(document.content)).paragraphs)
    assert first['redline']['suggested_text'] in text
    assert first['redline']['original_text'] not in text


def test_run_lock_and_restart_recovery(client):
    task = create_task(client)
    workspace_api.repo.start_run(task['id'])
    assert client.post('/api/tasks/' + task['id'] + '/review/stream', json={}).status_code == 409
    workspace_api.repo.interrupt_running()
    saved = workspace_api.repo.get(task['id'])
    assert saved['status'] == 'interrupted'
    review(client, saved)






def test_validation_prevents_fake_mode_inputs(client):
    base = {'title': '合同', 'text': '真实条款'}
    assert client.post('/api/tasks', json={**base, 'mode': 'checklist'}).status_code == 422
    assert client.post('/api/tasks', json={**base, 'mode': 'compare'}).status_code == 422
    assert client.post('/api/tasks', json={**base, 'text': ' '}).status_code == 422


def test_word_preserves_uploaded_table_layout_and_contains_native_comments(client):
    doc = Document()
    doc.sections[0].left_margin = Cm(3)
    doc.add_paragraph('企业采购合同', 'Title')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '金额'
    table.cell(0, 1).text = '1000元'
    doc.add_paragraph('第一条 付款')
    doc.add_paragraph('验收后30日付款。')
    data = BytesIO(); doc.save(data)
    imported = client.post('/api/contracts/import', files={'file': ('实际合同.docx', data.getvalue())}).json()
    task = create_task(client, title='企业采购合同', text=imported['text'], source_id=imported['source_id'])
    completed = review(client, task)
    run = workspace_api.repo.get(completed['run_id'])
    # Controlled revision fixture tests exporting actual applied text independently of the rules engine.
    issue = {'id': 'fixture', 'severity': 'warning', 'issue_title': '付款期限', 'clause_title': '第一条 付款',
             'risk_analysis': '付款期限需核对', 'legal_basis': '企业要求', 'negotiation_strategy': '按约定调整',
             'redline': {'original_text': '验收后30日付款。', 'suggested_text': '验收后60日付款。', 'change_type': 'replace'}}
    result = run['result']; result['issues'] = [issue]
    run = workspace_api.repo.update(run['id'], {'result': result})
    response = client.put('/api/runs/' + run['id'] + '/handling', json={'revision': run['revision'], 'applied': ['fixture'], 'comments': {'fixture': '业务确认付款期限'}})
    assert response.status_code == 200
    clean = client.get('/api/runs/' + run['id'] + '/export?scope=current')
    assert clean.headers['X-Word-Layout'] == 'preserved'
    exported = Document(BytesIO(clean.content))
    assert len(exported.tables) == 1
    assert exported.tables[0].cell(0, 1).text == '1000元'
    assert abs(exported.sections[0].left_margin.cm - 3) < .01
    assert '验收后60日付款。' in '\n'.join(p.text for p in exported.paragraphs)
    with ZipFile(BytesIO(clean.content)) as archive:
        assert 'word/comments.xml' in archive.namelist()
        assert '业务确认付款期限'.encode() in archive.read('word/comments.xml')
    tracked = client.get('/api/runs/' + run['id'] + '/export?scope=tracked')
    with ZipFile(BytesIO(tracked.content)) as archive:
        xml = archive.read('word/document.xml')
        assert b'<w:del ' in xml and b'<w:ins ' in xml and b'<w:delText' in xml
        assert b'commentRangeStart' in xml
        assert b'trackRevisions' in archive.read('word/settings.xml')


def test_comparison_model_assesses_only_actual_changes(client, monkeypatch):
    from backend import workspace_review
    seen = []
    async def fake_model(config, prompt):
        seen.append(prompt)
        return {'findings': [{'id': 'diff_1', 'severity': 'warning', 'explanation': '付款账期延长，请核对资金安排', 'basis': '交易要求', 'advice': '确认是否接受新账期', 'suggested_text': '验收后30日付款。'}]}
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: object())
    monkeypatch.setattr(workspace_review, 'model_json', fake_model)
    reference = '第一条 标的\n设备\n第二条 付款\n验收后30日付款。'
    task = create_task(client, mode='compare', text=reference.replace('30日', '60日'), reference_text=reference)
    result = review(client, task)['data']
    assert seen and '60日' in seen[0]
    assert result['issues'][0]['severity'] == 'warning'
    assert result['issues'][0]['redline']['original_text'] == '验收后60日付款。'




def test_unverified_checklist_evidence_cannot_be_marked_pass(client, monkeypatch):
    from backend import workspace_review
    async def fake_json(config, prompt):
        return {'findings': [{'id': 'balance', 'status': 'pass', 'evidence': '模型编造的原文', 'explanation': '通过'}]}
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: object())
    monkeypatch.setattr(workspace_review, 'model_json', fake_json)
    standard = checklist(client)
    task = create_task(client, mode='checklist', checklist_id=standard['id'])
    result = review(client, task)['data']
    finding = next(item for item in result['checklist_findings'] if item['id'] == 'balance')
    assert finding['status'] == 'needs_review'
    assert not finding['evidence']


def test_run_progress_is_persisted_and_list_responses_omit_contract_bodies(client):
    task = create_task(client)
    completed = review(client, task)
    run = client.get('/api/runs/' + completed['run_id']).json()
    assert len(run['progress_log']) == 5
    row = client.get('/api/tasks').json()['items'][0]
    assert row['id'] == task['id']
    assert 'text' not in row and 'run' not in row and 'result' not in row


def test_original_source_is_served_unchanged(client):
    data = BytesIO()
    document = Document()
    document.add_paragraph('原件预览测试')
    document.save(data)
    imported = client.post('/api/contracts/import', files={'file': ('original.docx', data.getvalue())}).json()
    source_id = imported['source_id']
    info = client.get('/api/sources/' + source_id)
    assert info.status_code == 200
    assert info.json()['extension'] == '.docx'
    assert 'text' not in info.json()
    file = client.get('/api/sources/' + source_id + '/file')
    assert file.status_code == 200
    assert file.content == data.getvalue()
    assert file.headers['content-type'].startswith('application/vnd.openxmlformats')
    assert file.headers['cache-control'] == 'no-store'
    assert client.get('/api/sources/' + '0' * 32 + '/file').status_code == 404


def test_formatted_source_edit_preserves_runs_tables_and_original(client):
    document = Document()
    document.sections[0].header.paragraphs[0].text = 'Header'
    p = document.add_paragraph()
    p.add_run('Amount: ').bold = True
    p.add_run('100 USD').italic = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Table text'
    document.add_paragraph('Special\nsoft break')
    data = BytesIO()
    document.save(data)
    imported = client.post('/api/contracts/import', files={'file': ('layout.docx', data.getvalue())}).json()
    old = imported['source_id']
    response = client.post('/api/sources/' + old + '/edit', json={'text': 'Amount: 200 USD\nTable revised\nSpecial\nsoft break'})
    assert response.status_code == 200, response.text
    edited = response.json()
    assert edited['source_id'] != old
    result = Document(BytesIO(client.get('/api/sources/' + edited['source_id'] + '/file').content))
    assert result.paragraphs[0].text == 'Amount: 200 USD'
    assert result.paragraphs[0].runs[0].bold
    assert result.paragraphs[0].runs[1].italic
    assert result.tables[0].cell(0, 0).text == 'Table revised'
    assert result.paragraphs[-1].text == 'Special\nsoft break'
    assert result.sections[0].header.paragraphs[0].text == 'Header'
    assert client.get('/api/sources/' + old + '/file').content == data.getvalue()
    assert client.post('/api/sources/' + old + '/edit', json={'text': 'Merged'}).status_code == 422


def test_current_document_download_does_not_require_review(client):
    task = create_task(client)
    response = client.get('/api/tasks/'+task['id']+'/document')
    assert response.status_code == 200
    assert task['text'].split('\n')[-1] in [p.text for p in Document(BytesIO(response.content)).paragraphs]
    review(client, task)
    latest = client.get('/api/tasks/'+task['id']).json()
    client.put('/api/tasks/'+task['id'], json={**latest, 'text': latest['text']+'\nNew manual edit'})
    current = client.get('/api/tasks/'+task['id']+'/document')
    assert 'New manual edit' in [p.text for p in Document(BytesIO(current.content)).paragraphs]
    snapshot = client.get('/api/tasks/'+task['id']+'/document?run='+latest['active_run'])
    assert 'New manual edit' not in [p.text for p in Document(BytesIO(snapshot.content)).paragraphs]


def test_adopted_document_preview_uses_same_formatted_document(client):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Amount: 100').bold = True
    buffer = BytesIO()
    document.save(buffer)
    imported = client.post('/api/contracts/import', files={'file': ('contract.docx', buffer.getvalue())}).json()
    task = create_task(client, text=imported['text'], source_id=imported['source_id'])
    review(client, task)
    run = client.get('/api/tasks/'+task['id']).json()['run']
    issue = run['result']['issues'][0]
    issue['redline'] = {'change_type': 'replace', 'original_text': 'Amount: 100', 'suggested_text': 'Amount: 200'}
    workspace_api.repo.update(run['id'], {'result': run['result'], 'applied': [issue['id']]})
    info = client.get('/api/sources/'+imported['source_id']+'?include_text=true&run='+run['id']).json()
    assert info['text'] == 'Amount: 200'
    file = client.get('/api/runs/'+run['id']+'/export?scope=current')
    preview = Document(BytesIO(file.content))
    assert preview.paragraphs[0].text == info['text']
    assert preview.paragraphs[0].runs[0].bold


def test_all_suggestions_preview_handles_sample(client):
    text = SAMPLE_CONTRACTS['business_cooperation_contract']['text']
    imported = client.post('/api/contracts/import', files={'file':('sample.txt',text.encode())}).json()
    task = create_task(client, text=text, source_id=imported['source_id'])
    run = review(client, task)
    info = client.get('/api/sources/'+imported['source_id']+'?include_text=true&run='+run['run_id']+'&scope=all')
    assert info.status_code == 200, info.text
    assert info.json()["preview_warnings"]
    response = client.get('/api/runs/'+run['run_id']+'/preview-file?scope=all')
    assert response.status_code == 200, response.text


def test_document_comparison_marks_replacements_insertions_deletions(client):
    response=client.post('/api/documents/compare',json={'original':'Title\nOld clause\nRemoved\nEnd','revised':'Title\nNew clause\nEnd\nAdded'})
    assert response.status_code==200
    data=response.json()
    assert len(data['original_map'])==4
    assert len(data['revised_map'])==4
    assert data['original_map'][0]==0
    assert data['original_map'][3]==2
    assert data['revised_map'][2]==3
    assert any(change['kind']=='insert' for change in data['changes'])
    assert any(change['original_start']==1 and change['original_end']==3 for change in data['changes'])
    same=client.post('/api/documents/compare',json={'original':'A\nB','revised':'A\nB'}).json()
    assert same['changes']==[]
    assert same['original_map']==[0,1]
