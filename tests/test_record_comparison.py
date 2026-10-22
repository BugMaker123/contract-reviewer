"""Formatted comparison, cumulative round baselines and explicit proposal scope."""
from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from playwright.sync_api import sync_playwright, expect
from backend.main import app
from backend import workspace_api
from backend.revisions import revised_text
from backend.store import now
from test_workflow_polish import completed, accept_payment, client
from test_ui_composition import live_server

HTML = '<h2 style="text-align:center">设备采购合同</h2><p><strong style="font-size:18pt;color:#c00000">第一条 设备交付。</strong></p><p>第二条 验收后30日付款。</p><table><tr><th>项目</th><th>说明</th></tr><tr><td>设备</td><td>一套</td></tr></table>'


def formatted_review(client):
    source = client.post('/api/contracts/paste', json={'title':'原格式对比验收','html':HTML}).json()
    item, run = completed(client, title='原格式对比验收', text=source['text'], source_id=source['source_id'])
    run = accept_payment(client, item, run)
    return item, run


def next_round(item, previous):
    run = workspace_api.repo.start_run(item['id'], baseline_kind='accepted',
                                       baseline_data={'text':revised_text(previous), 'source_id':None})
    run = workspace_api.repo.update(run['id'], {'status':'completed','completed_at':now(),
        'result':{'issues':[], 'segments':[], 'summary':{'verdict':'待人工复核','risk_level':'一般','compliance_score':80}}})
    workspace_api.repo.update(item['id'], {'status':'pending_review','active_run':run['id']})
    return run


def test_initial_comparison_keeps_prior_round_changes(client):
    item, first = formatted_review(client)
    second = next_round(item, first)
    base = '/api/tasks/'+item['id']+'/comparison'
    assert client.get(base, params={'after_run':second['id']}).json()['changes'] == []
    data = client.get(base, params={'before_run':'initial','after_run':second['id']}).json()
    assert data['before_number'] == 1 and data['after_number'] == 2
    assert len(data['changes']) == 1
    assert data['original_chars'] and data['revised_chars']
    assert '30日' in data['changes'][0]['original'] and '15日' in data['changes'][0]['revised']
    for side in ('original','revised'):
        response = client.get(base+'/document/'+side, params={'before_run':'initial','after_run':second['id']})
        assert response.status_code == 200, response.text
        doc = Document(BytesIO(response.content))
        assert doc.tables[0].cell(1, 0).text == '设备'
        assert any(r.bold and r.font.size.pt == 18 and str(r.font.color.rgb)=='C00000' for p in doc.paragraphs for r in p.runs)
        info = client.get(base+'/document/'+side, params={'before_run':'initial','after_run':second['id'],'info':True}).json()
        assert info['layout'] == 'preserved'
        assert info['text'] == data[side]


def test_pending_proposals_are_not_misreported_as_accepted(client):
    item, run = formatted_review(client)
    cleared = client.put('/api/runs/'+run['id']+'/handling', json={'revision':run['revision'],'applied':[]}).json()
    cleared = workspace_api.repo.get(run['id'], 'run')
    endpoint='/api/tasks/'+item['id']+'/comparison'
    current=client.get(endpoint).json()
    assert current['changes']==[] and current['pending_count']==1
    proposals=client.get(endpoint, params={'scope':'all'}).json()
    assert len(proposals['changes'])==1 and proposals['scope']=='all'
    assert workspace_api.repo.get(run['id'],'run') == cleared
    other, _=completed(client,title='另一任务')
    assert client.get('/api/tasks/'+other['id']+'/comparison/document/original',params={'after_run':run['id']}).status_code==422
    assert client.get(endpoint,params={'scope':'invalid'}).status_code==422


def test_browser_formatted_comparison_and_pending_scope(client):
    item, first=formatted_review(client)
    next_round(item,first)
    with live_server() as base, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1920,'height':1080})
        errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
        page.goto(base+'/records')
        page.get_by_role('button',name='前后对比',exact=True).click()
        dialog=page.get_by_role('dialog',name='审核前后对比')
        left=page.frame_locator('iframe[title="修改前合同原格式文档"]')
        right=page.frame_locator('iframe[title="修改后合同原格式文档"]')
        expect(left.locator('section.docx')).to_be_visible()
        expect(left.locator('table')).to_be_visible()
        expect(right.locator('table')).to_be_visible()
        expect(left.locator('mark.lexis-diff-token')).to_contain_text(['30'])
        expect(right.locator('mark.lexis-diff-token')).to_contain_text(['15'])
        styled=left.locator('article p').filter(has_text='第一条 设备交付。').locator('span').first
        assert styled.evaluate('(el)=>getComputedStyle(el).fontWeight') in ('700','bold')
        page.get_by_role('button',name='下一处差异').click()
        expect(right.locator('.comparison-current-change')).to_have_count(1)
        page.screenshot(path='artifacts/prototypes/record-comparison-formatted.png')
        page.get_by_label('仅看差异').check()
        expect(right.locator('article p').filter(has_text='第一条 设备交付。')).to_be_hidden()
        expect(right.locator('mark.lexis-diff-token')).to_be_visible()
        page.get_by_label('仅看差异').uncheck()
        expect(right.locator('table')).to_be_visible()
        page.get_by_label('比较基准').select_option('')
        expect(dialog.get_by_role('status')).to_contain_text('文字一致')
        expect(right.locator('article p').filter(has_text='验收后15日付款')).to_be_visible()
        page.get_by_role('button',name='关闭对比').click()
        browser.close()
        assert not errors, errors


def test_browser_zero_accepted_can_show_unaccepted_proposals(client):
    item, run = formatted_review(client)
    client.put('/api/runs/'+run['id']+'/handling',json={'revision':run['revision'],'applied':[]})
    with live_server() as base, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1600,'height':1000})
        page.goto(base+'/records')
        page.get_by_role('button',name='前后对比',exact=True).click()
        dialog=page.get_by_role('dialog',name='审核前后对比')
        expect(dialog.get_by_role('status')).to_contain_text('另有 1 项未采纳建议')
        expect(page.get_by_label('仅看差异')).to_be_disabled()
        right=page.frame_locator('iframe[title="修改后合同原格式文档"]')
        expect(right.locator('table')).to_be_visible()
        dialog.get_by_role('button',name='查看全部建议差异').click()
        expect(dialog.get_by_label('修改范围')).to_have_value('all')
        expect(right.locator('mark.lexis-diff-token')).to_contain_text(['15'])
        expect(dialog.get_by_text('建议预览 · 未自动采纳')).to_be_visible()
        assert workspace_api.repo.get(run['id'],'run')['applied'] == []
        page.get_by_label('仅看差异').check()
        page.get_by_label('修改范围').select_option('current')
        expect(page.get_by_label('仅看差异')).not_to_be_checked()
        expect(right.locator('table')).to_be_visible()
        page.set_viewport_size({'width':760,'height':900})
        assert dialog.evaluate('(el)=>el.scrollWidth<=el.clientWidth')
        expect(dialog.get_by_role('button',name='关闭对比')).to_be_visible()
        browser.close()
