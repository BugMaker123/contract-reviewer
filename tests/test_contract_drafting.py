from io import BytesIO
import re
import pytest
from fastapi.testclient import TestClient
from docx import Document
from backend.main import app
from backend import drafting_api, workspace_api

BRIEF={'title':'设备采购合同','contract_type':'采购合同','transaction':'采购一批设备，送货至甲方指定地址，验收后付款。','party_a':'甲方公司','party_b':'乙方公司','perspective':'甲方立场','requirements':'明确验收和质量责任'}
TEXT='设备采购合同\n甲方公司与乙方公司就设备采购达成协议。\n第一条 设备采购及交付安排。\n第二条 验收合格后按约支付价款。\n第三条 双方按约承担质量责任。'

@pytest.fixture
def client(monkeypatch):
    async def model(config,prompt):
        assert '甲方公司' in prompt and 'missing_information' in prompt
        return {'text':TEXT,'missing_information':['合同价款及付款期限']}
    monkeypatch.setattr(drafting_api,'model_json',model)
    with TestClient(app) as c: yield c

def generate(client):
    response=client.post('/api/contract-drafts/generate',json={'brief':BRIEF,'model_config_data':{'api_key':'test-secret'}})
    assert response.status_code==200,response.text
    return response.json()

def test_generate_edit_export_and_reopen(client):
    draft=generate(client)
    assert 'test-secret' not in str(workspace_api.repo.get(draft['id']))
    assert client.get('/api/contract-drafts').json()['total']==1
    edited=client.put('/api/contract-drafts/'+draft['id'],json={'revision':draft['revision'],'title':draft['title'],'text':TEXT+'\n补充条款：银行转账支付。'})
    assert edited.status_code==200
    assert client.put('/api/contract-drafts/'+draft['id'],json={'revision':draft['revision'],'title':draft['title'],'text':TEXT}).status_code==409
    saved=client.get('/api/contract-drafts/'+draft['id']).json()
    output=client.get('/api/contract-drafts/'+draft['id']+'/word')
    assert output.status_code==200
    exported = Document(BytesIO(output.content))
    assert '银行转账' in '\n'.join(p.text for p in exported.paragraphs)
    assert sum(p.text == draft['title'] for p in exported.paragraphs) == 1
    assert saved['text'].endswith('银行转账支付。')
    assert 'text' not in client.get('/api/contract-drafts').json()['items'][0]

def test_generation_failure_does_not_overwrite_saved_draft(client,monkeypatch):
    draft=generate(client)
    async def invalid(config,prompt): return {'text':'invalid'}
    monkeypatch.setattr(drafting_api,'model_json',invalid)
    assert client.post('/api/contract-drafts/generate',json={'brief':BRIEF}).status_code==502
    assert client.get('/api/contract-drafts/'+draft['id']).json()==draft
    assert client.get('/api/contract-drafts').json()['total']==1

def test_draft_records_search_and_pagination(client):
    for _ in range(3): generate(client)
    data=client.get('/api/contract-drafts?page=2&page_size=2').json()
    assert data['total']==3 and len(data['items'])==1
    assert client.get('/api/contract-drafts?q=不存在').json()['total']==0
    assert client.post('/api/contract-drafts/generate',json={'brief':{**BRIEF,'transaction':'短'}}).status_code==422

def test_drafting_browser_flow(client):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright,expect
    with live_server() as base,sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1600,'height':1000},accept_downloads=True)
        errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.goto(base+'/drafting')
        expect(page.get_by_role('heading',name='合同起草',exact=True)).to_be_visible()
        expect(page.get_by_label('搜索起草记录')).to_be_visible()
        expect(page.locator('.drafting-layout')).to_have_count(0)
        page.get_by_role('link',name='新建起草',exact=False).first.click()
        page.wait_for_url('**/drafting/new')
        expect(page.get_by_role('heading',name='新建合同起草',exact=True)).to_be_visible()
        expect(page.get_by_label('搜索起草记录')).to_have_count(0)
        expect(page.get_by_role('button',name='生成合同草稿',exact=True)).to_be_disabled()
        page.get_by_role('button',name='填入示例',exact=True).click()
        expect(page.get_by_label('交易内容与履行安排',exact=False)).to_have_value(re.compile('甲方向乙方采购'))
        expect(page.get_by_role('button',name='填入示例',exact=True)).to_be_disabled()
        focus=page.get_by_role('button',name='＋ 明确知识产权归属',exact=True)
        focus.click()
        expect(focus).to_have_attribute('aria-pressed','true')
        focus.click()
        expect(focus).to_have_attribute('aria-pressed','false')
        page.get_by_label('合同名称',exact=False).first.fill(BRIEF['title'])
        page.get_by_label('交易内容与履行安排',exact=False).fill(BRIEF['transaction'])
        page.get_by_label('甲方名称',exact=True).fill(BRIEF['party_a'])
        page.get_by_role('button',name='生成合同草稿',exact=True).click()
        page.wait_for_url('**/drafting/*/edit')
        edit_url=page.url
        expect(page.get_by_role('heading',name='编辑合同草稿',exact=True)).to_be_visible()
        expect(page.get_by_label('合同草稿正文')).to_have_value(TEXT)
        page.reload()
        expect(page.get_by_label('合同草稿正文')).to_have_value(TEXT)
        page.get_by_label('合同草稿正文').fill(TEXT+'\n补充条款：银行转账支付。')
        page.get_by_role('button',name='预览',exact=True).click()
        expect(page.get_by_label('合同草稿预览')).to_contain_text('银行转账支付')
        expect(page.get_by_text('有未保存修改',exact=True)).to_be_visible()
        page.get_by_role('button',name='编辑',exact=True).click()
        expect(page.get_by_label('合同草稿正文')).to_have_value(TEXT+'\n补充条款：银行转账支付。')
        with page.expect_download() as download:
            page.get_by_role('button',name='导出 Word').click()
        assert '银行转账' in '\n'.join(p.text for p in Document(download.value.path()).paragraphs)
        page.get_by_label('合同草稿正文').fill(TEXT+'\n补充条款：银行转账支付。\n签署后生效。')
        page.get_by_role('link',name='返回列表',exact=True).click()
        page.wait_for_url('**/drafting')
        expect(page.get_by_label('搜索起草记录')).to_be_visible()
        page.screenshot(path='artifacts/prototypes/contract-drafting-list.png')
        page.get_by_role('link',name='编辑草稿',exact=True).click()
        expect(page).to_have_url(edit_url)
        expect(page.get_by_label('合同草稿正文')).to_have_value(TEXT+'\n补充条款：银行转账支付。\n签署后生效。')
        page.locator('.portal-main').evaluate('(el)=>{el.scrollTop=0}')
        page.screenshot(path='artifacts/prototypes/contract-drafting.png')
        page.set_viewport_size({'width':1000,'height':900})
        brief_box=page.locator('.drafting-brief').bounding_box()
        output_box=page.locator('.drafting-output').bounding_box()
        assert output_box['y'] >= brief_box['y']+brief_box['height']
        assert page.locator('.drafting-page').evaluate('(el)=>el.scrollWidth<=el.clientWidth')
        page.set_viewport_size({'width':1600,'height':1000})
        page.get_by_role('button',name='进入合同审查',exact=False).click()
        page.wait_for_url('**/workbench?task=*')
        task_id=page.url.split('task=')[1]
        task=workspace_api.repo.get(task_id,'task')
        assert task['text'].endswith('签署后生效。') and task['status']=='draft'
        assert not errors,errors
        browser.close()
