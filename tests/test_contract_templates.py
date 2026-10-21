from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from backend.main import app


def test_template_lifecycle_and_original_download():
    with TestClient(app) as client:
        group = client.post('/api/template-groups', json={'name': '采购协议'}).json()
        group2 = client.post('/api/template-groups', json={'name': '经销协议'}).json()
        assert client.post('/api/template-groups', json={'name': ' 采购协议 '}).status_code == 409
        assert client.post('/api/template-groups', json={'name': '   '}).status_code == 422
        doc = Document()
        doc.add_paragraph('甲方：____；付款期限为30日。')
        data = BytesIO()
        doc.save(data)
        content = data.getvalue()
        source = client.post('/api/contracts/import', files={'file': ('协议.docx', content)}).json()
        body = {'name': '年度采购模板', 'description': '2026修订版', 'group_id': group['id'], 'source_id': source['source_id']}
        response = client.post('/api/contract-templates', json=body)
        assert response.status_code == 200, response.text
        template = response.json()
        url = '/api/contract-templates/' + template['id']
        assert client.get(url).json()['filename'] == '协议.docx'
        download = client.get(url + '/download')
        assert download.content == content
        assert 'attachment' in download.headers['content-disposition']
        assert client.get('/api/sources/' + source['source_id'] + '?include_text=true').json()['text']
        assert client.get('/api/contract-templates?q=2026').json()['total'] == 1
        assert client.get('/api/contract-templates', params={'group_id': group2['id']}).json()['total'] == 0
        assert client.delete('/api/template-groups/' + group['id'] + '?revision=1').status_code == 409
        assert client.put(url, json={**body, 'revision': 999}).status_code == 409
        updated = client.put(url, json={**body, 'name': '新版', 'group_id': group2['id'], 'revision': 1}).json()
        assert updated['revision'] == 2
        assert client.delete('/api/template-groups/' + group['id'] + '?revision=1').status_code == 200
        assert client.post('/api/contract-templates', json=body).status_code == 404
        assert client.delete(url + '?revision=1').status_code == 409
        assert client.delete(url + '?revision=2').status_code == 200
        assert client.get(url + '/download').status_code == 404
        assert client.delete('/api/template-groups/' + group2['id'] + '?revision=1').status_code == 200


def test_group_update_and_template_validation():
    with TestClient(app) as client:
        group = client.post('/api/template-groups', json={'name': '默认'}).json()
        url = '/api/template-groups/' + group['id']
        assert client.put(url, json={'name': '修改'}).status_code == 422
        assert client.put(url, json={'name': '修改', 'revision': 1}).json()['name'] == '修改'
        assert client.put(url, json={'name': '旧页', 'revision': 1}).status_code == 409
        assert client.post('/api/contract-templates', json={'name': '无效原件', 'group_id': group['id'], 'source_id': 'missing'}).status_code == 404
        assert client.get('/api/contract-templates?page=0').status_code == 422


def test_template_browser_upload_preview_and_delete():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    document = Document()
    document.add_paragraph('模板正文预览测试')
    output = BytesIO()
    document.save(output)
    with live_server() as base, sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={'width': 1600, 'height': 1000})
        errors = []
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.goto(base + '/templates')
        expect(page.get_by_role('heading', name='合同模板', exact=True)).to_be_visible()
        page.get_by_role('button', name='新建分组', exact=True).click()
        dialog = page.get_by_role('dialog')
        dialog.get_by_label('分组名称', exact=True).fill('采购模板组')
        dialog.get_by_role('button', name='保存', exact=True).click()
        expect(dialog).to_have_count(0)
        page.locator('.catalog-heading-actions').get_by_role('button', name='新增模板', exact=True).click()
        dialog.locator('input[type=file]').set_input_files({'name': 'sample.docx', 'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'buffer': output.getvalue()})
        dialog.get_by_label('模板名称', exact=True).fill('年度协议')
        dialog.get_by_label('所属分组',exact=True).select_option(label='采购模板组')
        dialog.get_by_role('button', name='创建模板', exact=True).click()
        expect(dialog).to_have_count(0, timeout=15000)
        row = page.get_by_role('row').filter(has_text='年度协议')
        row.get_by_role('button', name='更多操作 年度协议', exact=True).click()
        page.get_by_role('menuitem', name='预览原件', exact=True).click()
        expect(page.frame_locator('iframe').locator('#pages')).to_contain_text('模板正文预览测试', timeout=15000)
        page.get_by_role('button', name='关闭', exact=True).click()
        row.get_by_role('button', name='更多操作 年度协议', exact=True).click()
        with page.expect_download() as download:
            page.get_by_role('menuitem', name='下载原件', exact=True).click()
        assert download.value.suggested_filename == 'sample.docx'
        row.get_by_role('button', name='更多操作 年度协议', exact=True).click()
        page.get_by_role('menuitem', name='删除模板', exact=True).click()
        page.get_by_role('button', name='确认删除', exact=True).click()
        expect(row).to_have_count(0)
        page.reload()
        expect(page.get_by_text('采购模板组', exact=False).first).to_be_visible()
        assert not errors
        browser.close()


def test_template_list_status_filter_and_field_summary():
    with TestClient(app) as client:
        group=client.post('/api/template-groups',json={'name':'采购'}).json()
        doc=Document();doc.add_paragraph('甲方：____')
        data=BytesIO();doc.save(data)
        source=client.post('/api/contracts/import',files={'file':('source.docx',data.getvalue())}).json()['source_id']
        template=client.post('/api/contract-templates',json={'name':'可配置模板','group_id':group['id'],'source_id':source}).json()
        assert client.get('/api/contract-templates?status=unconfigured').json()['total']==1
        fields=[{'id':'f1','name':'甲方','block_id':'b0','start':3,'end':7,'confirmed':False}]
        base='/api/contract-templates/'+template['id']
        client.put(base+'/review',json={'revision':1,'source_id':source,'fields':fields})
        summary=client.get('/api/contract-templates?status=reviewing').json()['items'][0]
        assert summary['field_count']==1 and summary['confirmed_count']==0
        assert 'review_draft' not in summary
        fields[0]['confirmed']=True
        client.put(base+'/review',json={'revision':2,'source_id':source,'fields':fields})
        assert client.get('/api/contract-templates?status=ready').json()['total']==1
        client.post(base+'/publish',json={'revision':3})
        assert client.get('/api/contract-templates?status=published').json()['total']==1
        assert client.get('/api/contract-templates?status=unconfigured').json()['total']==0
        assert client.get('/api/contract-templates?status=unknown').status_code==422


def test_browser_add_template_without_precreating_group():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright,expect
    doc=Document();doc.add_paragraph('模板正文')
    data=BytesIO();doc.save(data)
    with live_server() as address,sync_playwright() as p:
        browser=p.chromium.launch();page=browser.new_page(viewport={'width':1600,'height':1000})
        page.goto(address+'/templates')
        page.locator('.catalog-heading-actions').get_by_role('button',name='新增模板',exact=True).click()
        dialog=page.get_by_role('dialog')
        dialog.get_by_label('模板名称',exact=True).fill('独立模板')
        dialog.locator('input[type=file]').set_input_files({'name':'source.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':data.getvalue()})
        dialog.get_by_role('button',name='创建模板',exact=True).click()
        expect(dialog).to_have_count(0,timeout=15000)
        row=page.get_by_role('row').filter(has_text='独立模板')
        expect(row).to_contain_text('未分组')
        expect(row.locator('.template-state')).to_have_text('待配置')
        page.get_by_label('筛选配置状态',exact=True).select_option('published')
        expect(row).to_have_count(0)
        page.get_by_label('筛选配置状态',exact=True).select_option('all')
        expect(row).to_be_visible()
        page.set_viewport_size({'width':390,'height':844})
        assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
        browser.close()
