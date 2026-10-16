from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.main import app
from backend import workspace_api, workspace_review
from backend.word_export import paragraphs, visible_paragraph_text


@pytest.fixture
def revision_client(monkeypatch):
    monkeypatch.setattr(workspace_review.graph, 'get_llm_client', lambda config: None)
    with TestClient(app) as client:
        yield client


def seed_review(client, with_source=False):
    first, second = '1. 验收后付款。', '2. 违约责任另行协商。'
    text = first + '\n' + second
    source_id = None
    if with_source:
        document = Document()
        document.add_paragraph(first)
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = second
        buffer = BytesIO()
        document.save(buffer)
        imported = client.post('/api/contracts/import', files={'file': ('review.docx', buffer.getvalue())}).json()
        text, source_id = imported['text'], imported['source_id']
    task = client.post('/api/tasks', json={'title': '修订模式合同', 'text': text, 'mode': 'quick', 'source_id': source_id}).json()
    client.post('/api/tasks/'+task['id']+'/review/stream', json={})
    task = client.get('/api/tasks/'+task['id']).json()
    run = task['run']
    issues = []
    for identifier, original, suggested in [('payment', first, '1. 验收合格后30日付款。'), ('liability', second, '2. 累计违约赔偿不超过合同总价。')]:
        issues.append({'id': identifier, 'segment_id': identifier, 'clause_title': '付款条款' if identifier=='payment' else '责任条款', 'severity': 'warning', 'category': '商务', 'issue_title': '约定不明确', 'risk_analysis': '需要明确条件', 'legal_basis': '', 'negotiation_strategy': '', 'evidence': original, 'redline': {'original_text': original, 'suggested_text': suggested, 'change_type': 'replace'}})
    result = {**run['result'], 'issues': issues}
    run = workspace_api.repo.update(run['id'], {'result': result})
    return task, run


def native_text(content):
    document = Document(BytesIO(content))
    return '\n'.join(visible_paragraph_text(p._p) for p in paragraphs(document)), document


@pytest.mark.parametrize('with_source', [False, True])
def test_native_revision_document_accepts_and_rejects_independently(revision_client, with_source):
    client = revision_client
    task, run = seed_review(client, with_source)
    markup = client.get('/api/runs/'+run['id']+'/review-document')
    assert markup.status_code == 200
    with ZipFile(BytesIO(markup.content)) as archive:
        xml = archive.read('word/document.xml')
        assert b'<w:del ' in xml and b'<w:ins ' in xml
        assert b'trackRevisions' in archive.read('word/settings.xml')
    info = client.get('/api/runs/'+run['id']+'/review-document/info').json()
    assert len(info['review_items']) == 2 and not info['revision_warnings']
    assert info['layout'] == ('preserved' if with_source else 'reconstructed')
    if with_source:
        _, pending_document = native_text(markup.content)
        cell = pending_document.tables[0].cell(0, 0)
        cell_text = '\n'.join(visible_paragraph_text(p._p) for p in cell.paragraphs)
        assert cell_text == '2. 累计违约赔偿不超过合同总价。'
        assert visible_paragraph_text(pending_document.paragraphs[0]._p) == '1. 验收合格后30日付款。'
    accepted = client.put('/api/runs/'+run['id']+'/handling', json={'revision':run['revision'], 'applied':['payment'], 'decisions':{'liability':'rejected'}}).json()
    markup = client.get('/api/runs/'+run['id']+'/review-document')
    text, document = native_text(markup.content)
    assert '30日付款' in text and '另行协商' in text and '累计违约' not in text
    assert not document.element.xpath('//w:ins') and not document.element.xpath('//w:del')
    assert len(document.tables) == (1 if with_source else 0)
    original, _ = native_text(client.get('/api/runs/'+run['id']+'/review-document?display=original').content)
    assert '30日付款' not in original and '验收后付款' in original
    assert client.get('/api/tasks/'+task['id']).json()['text'] == run['input']['text']


def test_partial_acceptance_leaves_only_pending_edits_tracked(revision_client):
    client = revision_client
    _, run = seed_review(client, True)
    client.put('/api/runs/'+run['id']+'/handling', json={'revision':run['revision'], 'applied':['payment']})
    markup = client.get('/api/runs/'+run['id']+'/review-document')
    text, document = native_text(markup.content)
    assert '30日付款' in text and '累计违约赔偿' in text
    deleted = ''.join(node.text or '' for node in document.element.iter(qn('w:delText')))
    assert '另行协商' in deleted and '付款' not in deleted
    current, _ = native_text(client.get('/api/runs/'+run['id']+'/review-document?display=current').content)
    assert '30日付款' in current and '另行协商' in current


def test_regenerate_preserves_versions_resets_acceptance_and_restores(revision_client, monkeypatch):
    client = revision_client
    _, run = seed_review(client)
    run = client.put('/api/runs/'+run['id']+'/handling', json={'revision':run['revision'], 'applied':['payment']}).json()
    before = deepcopy(run['result']['issues'][0]['redline'])
    async def generate(config, prompt):
        assert '措辞温和' in prompt and '付款' in prompt
        return {'suggested_text':'1. 验收合格且收到发票后30日内付款。', 'explanation':'明确开票与付款期限'}
    monkeypatch.setattr(workspace_review, 'model_json', generate)
    result = client.post('/api/runs/'+run['id']+'/issues/payment/regenerate', json={'revision':run['revision'], 'instruction':'措辞温和'})
    assert result.status_code == 200, result.text
    generated = result.json()
    assert generated['applied'] == [] and generated['decisions']['payment'] == 'pending'
    assert generated['revision_history']['payment'][0]['redline'] == before
    assert len(generated['revision_history']['payment']) == 2
    assert generated['result']['issues'][1] == run['result']['issues'][1]
    restored = client.post('/api/runs/'+run['id']+'/issues/payment/revision/restore', json={'revision':generated['revision'], 'version':1}).json()
    assert restored['result']['issues'][0]['redline'] == before
    assert len(restored['revision_history']['payment']) == 3


@pytest.mark.parametrize('reply', [None, {'suggested_text':''}, {'suggested_text':'1. 验收后付款。'}, {'suggested_text':'1. 验收合格后30日付款。'}])
def test_failed_or_unchanged_ai_keeps_previous_revision(revision_client, monkeypatch, reply):
    client = revision_client
    _, run = seed_review(client)
    async def generate(config, prompt):
        if reply is None: raise ConnectionError('offline')
        return reply
    monkeypatch.setattr(workspace_review, 'model_json', generate)
    response = client.post('/api/runs/'+run['id']+'/issues/payment/regenerate', json={'revision':run['revision']})
    assert response.status_code == 502
    assert client.get('/api/runs/'+run['id']).json() == run


def test_manual_candidate_and_optimistic_conflict(revision_client):
    client = revision_client
    task, run = seed_review(client)
    url = '/api/runs/'+run['id']+'/issues/payment/revision'
    edited = client.put(url, json={'revision':run['revision'], 'suggested_text':'1. 双方确认验收结果后15日付款。'})
    assert edited.status_code == 200, edited.text
    assert edited.json()['decisions']['payment'] == 'pending'
    assert client.put(url, json={'revision':run['revision'], 'suggested_text':'过期修改'}).status_code == 409
    latest = client.get('/api/tasks/'+task['id']).json()
    client.put('/api/tasks/'+task['id'], json={**latest, 'text':latest['text']+'\n修改后的新条款', 'checklist_id':None})
    assert client.put(url, json={'revision':edited.json()['revision'], 'suggested_text':'过期结果'}).status_code == 409


def test_generation_cannot_overwrite_concurrent_handling(revision_client, monkeypatch):
    client = revision_client
    _, run = seed_review(client)
    async def generate(config, prompt):
        workspace_api.repo.update(run['id'], {'comments': {'liability':'其他页面的备注'}}, run['revision'])
        return {'suggested_text':'1. 验收合格后15日付款。'}
    monkeypatch.setattr(workspace_review, 'model_json', generate)
    response = client.post('/api/runs/'+run['id']+'/issues/payment/regenerate', json={'revision':run['revision']})
    assert response.status_code == 409
    saved = client.get('/api/runs/'+run['id']).json()
    assert saved['comments']['liability'] == '其他页面的备注'
    assert saved['result']['issues'] == run['result']['issues']


def test_word_style_revision_browser(revision_client, monkeypatch):
    import socket
    import threading
    import uvicorn
    from playwright.sync_api import sync_playwright, expect
    client = revision_client
    task, run = seed_review(client, True)
    async def generate(config, prompt):
        return {'suggested_text':'1. 验收合格且收到发票后15日付款。', 'explanation':'按本次要求调整开票和付款期限'}
    monkeypatch.setattr(workspace_review, 'model_json', generate)
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets':[sock]}, daemon=True)
    thread.start()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={'width':1600,'height':1050})
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(f'http://127.0.0.1:{sock.getsockname()[1]}/workbench?task='+task['id'])
            expect(page.get_by_test_id('tracked-document-pane')).to_be_visible()
            expect(page.get_by_test_id('original-document-pane')).to_have_count(0)
            preview = page.frame_locator('iframe[title="合同修订文档"]')
            expect(preview.locator('article ins')).not_to_have_count(0)
            expect(preview.locator('article del')).not_to_have_count(0)
            assert preview.locator('article table').count() == 1
            expect(preview.locator('#review-selection')).to_be_visible()
            assert preview.locator('#review-selection').count() == 1
            assert preview.locator('#review-callout').count() == 1
            assert page.locator('#review-summary-panel').bounding_box()['width'] >= 350
            # Selection labels belong to the single overlay, never individual paragraphs.
            assert preview.locator('.review-active').evaluate_all("els => els.every(el => ['none', 'normal'].includes(getComputedStyle(el, '::after').content))")
            page.screenshot(path='artifacts/word-review-markup.png')
            payment = page.locator('#card-payment')
            assert [label.strip() for label in payment.locator('.revision-primary-actions button').all_text_contents()] == ['采纳', '拒绝', '重新生成']
            expect(payment.get_by_role('button', name='手动调整修订', exact=True)).not_to_be_visible()
            expect(page.locator('.prototype-navbar')).to_be_visible()
            expect(page.locator('.prototype-view-switch')).to_be_visible()
            expect(page.locator('.wb-toolbar')).to_have_count(0)
            page.keyboard.press('Alt+ArrowDown')
            page.keyboard.press('Alt+ArrowUp')
            preview.locator('body').evaluate("() => { window.reviewInstance = 'preserved'; window.originalReviewArticle = document.querySelector('article'); }")
            original_preview_url = page.locator('iframe[title="合同修订文档"]').get_attribute('src')
            payment.get_by_role('button', name='采纳', exact=True).click()
            expect(payment.get_by_text('处理状态：已采纳', exact=True)).to_be_visible()
            expect(preview.locator('.review-applied')).not_to_have_count(0)
            assert preview.locator('body').evaluate("() => window.reviewInstance === 'preserved' && window.originalReviewArticle === document.querySelector('article')")
            assert page.locator('iframe[title="合同修订文档"]').get_attribute('src') == original_preview_url
            expect(payment).to_have_attribute('data-decision', 'applied')
            expect(preview.locator('article')).to_contain_text('30日付款')
            expect(preview.locator('article del')).not_to_contain_text('付款')
            page.get_by_role('tab', name='当前稿', exact=True).click()
            expect(preview.locator('article')).to_contain_text('30日付款')
            expect(preview.locator('article')).to_contain_text('另行协商')
            expect(preview.locator('article ins')).to_have_count(0)
            page.get_by_role('tab', name='所有标记', exact=True).click()
            payment.locator('.revision-generation-options > summary').click()
            payment.get_by_label('付款条款AI调整要求', exact=True).fill('补充开票条件，付款期限调整为15日')
            payment.get_by_role('button', name='重新生成', exact=True).click()
            expect(payment.get_by_text('处理状态：待处理', exact=True)).to_be_visible()
            expect(preview.locator('article')).to_contain_text('15日付款')
            expect(payment.get_by_text('第 2 版', exact=True)).to_be_visible()
            payment.locator('.revision-more-actions>summary').click()
            payment.get_by_role('button', name='手动调整修订', exact=True).click()
            payment.get_by_label('付款条款修订文本', exact=True).fill('1. 验收合格后20日付款。')
            payment.get_by_role('button', name='保存修订建议', exact=True).click()
            expect(preview.locator('article')).to_contain_text('20日付款')
            payment.locator('.revision-history>summary').click()
            payment.locator('.revision-history article').filter(has_text='第 1 版').get_by_role('button', name='恢复此版', exact=True).click()
            expect(preview.locator('article')).to_contain_text('30日付款')
            page.reload()
            expect(page.get_by_test_id('tracked-document-pane')).to_be_visible()
            expect(preview.locator('article')).to_contain_text('30日付款')
            page.locator('#card-liability').get_by_role('button', name='拒绝', exact=True).click()
            expect(preview.locator('article')).to_contain_text('另行协商')
            expect(preview.locator('article')).not_to_contain_text('累计违约赔偿')
            with page.expect_download() as download:
                page.get_by_role('button', name='导出', exact=True).click()
            assert download.value.suggested_filename.endswith('-修订稿.docx')
            payment.get_by_role('button', name='采纳', exact=True).click()
            expect(payment.get_by_text('处理状态：已采纳', exact=True)).to_be_visible()
            expect(preview.locator('article ins')).to_have_count(0)
            current = client.get('/api/runs/'+run['id']).json()
            expect(page.get_by_test_id('tracked-document-pane')).to_have_attribute('data-rendered-revision', str(current['revision']))
            page.screenshot(path='artifacts/word-review-accepted.png')
            assert not errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=5)


def test_changed_review_mode_blocks_regeneration(revision_client, monkeypatch):
    client = revision_client
    task, run = seed_review(client)
    current = client.get('/api/tasks/'+task['id']).json()
    response = client.put('/api/tasks/'+task['id'], json={**current, 'mode':'deep'})
    assert response.status_code == 200
    async def must_not_generate(config, prompt):
        raise AssertionError('A stale review must not call the model')
    monkeypatch.setattr(workspace_review, 'model_json', must_not_generate)
    response = client.post('/api/runs/'+run['id']+'/issues/payment/regenerate', json={'revision':run['revision']})
    assert response.status_code == 409
    assert client.get('/api/runs/'+run['id']).json() == run


def test_adjacent_revisions_keep_soft_breaks_cell_positions_and_font(revision_client):
    from docx.shared import Pt
    client = revision_client
    original = Document()
    paragraph = original.add_paragraph()
    run = paragraph.add_run('Before\nsoft line')
    run.font.size = Pt(15)
    original.add_table(rows=1, cols=1).cell(0, 0).text = 'Second clause'
    buffer = BytesIO()
    original.save(buffer)
    imported = client.post('/api/contracts/import', files={'file':('soft.docx',buffer.getvalue())}).json()
    task = client.post('/api/tasks', json={'title':'Soft breaks','text':imported['text'],'source_id':imported['source_id'],'mode':'quick'}).json()
    client.post('/api/tasks/'+task['id']+'/review/stream',json={})
    run = client.get('/api/tasks/'+task['id']).json()['run']
    issues = []
    for identifier, old, new in [('one','Before\nsoft line','Before\nupdated soft line'),('two','Second clause','Second revised')]:
        issues.append({'id':identifier,'evidence':old,'issue_title':'Revision','redline':{'original_text':old,'suggested_text':new,'change_type':'replace'}})
    workspace_api.repo.update(run['id'], {'result':{**run['result'],'issues':issues}})
    response = client.get('/api/runs/'+run['id']+'/review-document')
    assert response.status_code == 200
    _, revised = native_text(response.content)
    assert visible_paragraph_text(revised.paragraphs[0]._p) == 'Before\nupdated soft line'
    assert visible_paragraph_text(revised.tables[0].cell(0,0).paragraphs[0]._p) == 'Second revised'
    assert revised.paragraphs[0].runs[0].font.size.pt == 15


def test_reference_layout_with_multiline_contract(revision_client):
    import socket
    import threading
    import uvicorn
    from playwright.sync_api import sync_playwright, expect

    client = revision_client
    document = Document()
    document.add_heading('年度采购协议', 0)
    document.add_paragraph('甲方（供方）：示例医药有限公司')
    document.add_paragraph('乙方（需方）：示例采购有限公司')
    document.add_heading('一、前提条件', 1)
    original = [
        '1.1 双方应具备本协议项下业务所需的合法资质，并提供营业执照及相关授权资料。',
        '双方保证提供的证件、资料真实、合法、有效，并及时通知对方资质变更情况。',
    ]
    for text in original:
        document.add_paragraph(text)
    for index in range(2, 20):
        document.add_paragraph(f'1.{index} 甲乙双方应按照本协议约定履行合同义务，及时核对采购数量、产品质量、发票及付款信息。发生争议时应及时书面通知对方，并协商确定处理期限和责任承担方式。')
    document.add_paragraph(original[0])  # A repeated fragment must not join the selected range.
    buffer = BytesIO()
    document.save(buffer)
    source = client.post('/api/contracts/import', files={'file': ('layout.docx', buffer.getvalue())}).json()
    task = client.post('/api/tasks', json={'title': '年度采购协议 · 布局验证', 'text': source['text'], 'mode': 'quick', 'source_id': source['source_id']}).json()
    client.post('/api/tasks/'+task['id']+'/review/stream', json={})
    run = client.get('/api/tasks/'+task['id']).json()['run']
    issues = []
    for identifier, evidence in [('qualification', '\n'.join(original)), ('payment', '1.2 甲乙双方应按照本协议约定履行合同义务，及时核对采购数量、产品质量、发票及付款信息。发生争议时应及时书面通知对方，并协商确定处理期限和责任承担方式。')]:
        issues.append({'id': identifier, 'segment_id': identifier, 'clause_title': '资质与通知' if identifier == 'qualification' else '付款条件', 'severity': 'critical', 'category': '商务', 'issue_title': '履约条件和责任承担需要明确', 'risk_analysis': '资质变更通知期限及责任范围未明确，建议约定书面通知时间并限定直接损失。', 'evidence': evidence, 'redline': {'original_text': evidence, 'suggested_text': evidence+' 资质变更应在五个工作日内书面通知，赔偿限于直接实际损失。', 'change_type': 'replace'}})
    workspace_api.repo.update(run['id'], {'result': {**run['result'], 'issues': list(reversed(issues))}})
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={'width': 1920, 'height': 1000})
            page.goto(f'http://127.0.0.1:{sock.getsockname()[1]}/workbench?task='+task['id'])
            preview = page.frame_locator('iframe[title="合同修订文档"]')
            expect(preview.locator('article')).to_contain_text('资质变更')
            expect(page.locator('.prototype-risk-card').first).to_have_attribute('id', 'card-qualification')
            page.locator('#card-qualification .prototype-locate').click()
            expect(preview.locator('#review-callout')).to_be_visible()
            expect(preview.locator('#review-selection')).to_be_visible()
            expect(preview.locator('.review-active')).to_have_count(2)
            assert preview.locator('#review-selection').count() == 1
            assert page.locator('#review-summary-panel').bounding_box()['width'] >= 390
            buttons = page.locator('#card-qualification .revision-primary-actions button')
            positions = [buttons.nth(index).bounding_box() for index in range(3)]
            assert max(box['y'] for box in positions) - min(box['y'] for box in positions) < 2
            assert positions[2]['x'] < positions[0]['x'] < positions[1]['x']
            page.wait_for_timeout(400)
            page.screenshot(path='artifacts/review-reference-layout.png')
            page.locator('#card-payment .prototype-locate').click()
            expect(preview.locator('.review-active')).to_have_count(1)
            expect(page.locator('#card-payment')).to_have_class(__import__('re').compile('ring-2'))
            page.wait_for_timeout(400)
            preview.locator('#viewport').evaluate('(el) => { el.scrollTop = 120; }')
            page.wait_for_timeout(250)
            preview.locator('body').evaluate("() => { window.reviewInstance = 'stable'; }")
            scroll_before = preview.locator('#viewport').evaluate('(el) => el.scrollTop')
            payment_card = page.locator('#card-payment')
            payment_card.get_by_role('button', name='采纳', exact=True).click()
            expect(preview.locator('.review-applied')).not_to_have_count(0)
            assert abs(preview.locator('#viewport').evaluate('(el) => el.scrollTop') - scroll_before) < 3
            expect(payment_card).to_have_attribute('data-decision', 'applied')
            page.screenshot(path='artifacts/review-applied-green.png')
            page.locator('#card-qualification').get_by_role('button', name='拒绝', exact=True).click()
            expect(preview.locator('.review-rejected')).not_to_have_count(0)
            assert preview.locator('body').evaluate("() => window.reviewInstance") == 'stable'
            assert abs(preview.locator('#viewport').evaluate('(el) => el.scrollTop') - scroll_before) < 3
            page.wait_for_timeout(400)
            preview.locator('#viewport').evaluate('(el) => { el.scrollTop = el.scrollHeight; }')
            expect(preview.locator('#review-callout')).not_to_be_visible()
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=5)
