"""Delivery requires manual completion; preview remains available while reviewing."""
from io import BytesIO
from docx import Document
from test_tracked_revisions import revision_client, seed_review


def accept_and_complete(client, task, run):
    ids = [issue['id'] for issue in run['result']['issues']]
    response = client.put('/api/runs/' + run['id'] + '/handling', json={
        'revision': run['revision'], 'applied': ids, 'decisions': {id: 'applied' for id in ids}, 'comments': {}})
    assert response.status_code == 200, response.text
    current = client.get('/api/tasks/' + task['id']).json()
    result = client.post('/api/tasks/' + task['id'] + '/complete', json={'revision': current['revision'], 'archived': False})
    assert result.status_code == 200, result.text
    return client.get('/api/tasks/' + task['id']).json()


def test_exports_require_current_manual_completion(revision_client):
    client = revision_client
    task, run = seed_review(client, True)
    urls = ['/api/runs/' + run['id'] + suffix for suffix in (
        '/export?scope=current', '/audit', '/preview-file', '/review-document')]
    urls.append('/api/tasks/' + task['id'] + '/document')
    for url in urls:
        assert client.get(url).status_code == 409, url
    assert client.get('/api/runs/' + run['id'] + '/export?scope=report').status_code == 200
    preview = client.get('/api/runs/' + run['id'] + '/review-document/preview')
    assert preview.status_code == 200
    assert preview.headers['content-disposition'].startswith('inline;')
    assert Document(BytesIO(preview.content)).paragraphs
    current = accept_and_complete(client, task, run)
    for url in urls:
        assert client.get(url).status_code == 200, url
    saved = current['run']
    response = client.put('/api/runs/' + run['id'] + '/handling', json={
        'revision': saved['revision'], 'applied': saved['applied'], 'decisions': saved['decisions'], 'comments': {'payment': '需要重新确认'}})
    assert response.status_code == 200
    for url in urls:
        assert client.get(url).status_code == 409, url


def test_completed_exports_remain_locked_after_input_change(revision_client):
    client = revision_client
    task, run = seed_review(client)
    current = accept_and_complete(client, task, run)
    changed = client.put('/api/tasks/' + task['id'], json={**current, 'text': current['text'] + '\n追加条款', 'source_id': None, 'checklist_id': None})
    assert changed.status_code == 200
    assert client.get('/api/runs/' + run['id'] + '/export').status_code == 409


def test_browser_compact_export_only_after_completion(revision_client):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    task, _ = seed_review(revision_client, True)
    with live_server() as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(base + '/workbench?task=' + task['id'])
        expect(page.get_by_test_id('tracked-document-pane')).to_be_visible()
        expect(page.get_by_role('button', name='编辑正文', exact=True)).to_have_count(0)
        assert page.locator('.prototype-nav-right').get_by_role('button').last.inner_text().strip() == '导出'
        export = page.get_by_role('button', name='导出', exact=True)
        menu = page.get_by_role('menu', name='导出审核文件')
        export.click()
        expect(page.get_by_text('请先完成审核，再导出文件', exact=True)).to_be_visible()
        expect(menu).to_have_count(0)
        page.get_by_role('button', name='审查意见书', exact=True).click()
        expect(page.get_by_role('button', name='下载 Word', exact=True)).to_be_enabled()
        with page.expect_download() as report:
            page.get_by_role('button', name='下载 Word', exact=True).click()
        assert Document(report.value.path()).paragraphs
        expect(page.get_by_role('button', name='打印 / 存为 PDF', exact=True)).to_be_enabled()
        page.get_by_role('button', name='关闭意见书', exact=True).click()
        page.get_by_role('button', name='一键全部采纳', exact=True).click()
        expect(page.get_by_role('button', name='一键全部采纳', exact=True)).to_be_disabled()
        page.get_by_role('button', name='完成审核', exact=True).click()
        expect(page.get_by_role('button', name='已完成审核', exact=True)).to_be_visible()
        assert '/workbench?task=' + task['id'] in page.url
        for title in ['修订稿']:
            export.click()
            expect(menu.get_by_role('menuitem')).to_have_count(1)
            expect(menu.get_by_text('审查意见书', exact=True)).to_have_count(0)
            box = menu.bounding_box()
            assert box['width'] <= 250 and box['height'] < 160
            page.screenshot(path='artifacts/compact-completed-export.png')
            with page.expect_download() as download:
                menu.get_by_text(title, exact=True).click()
            assert download.value.failure() is None
            doc = Document(download.value.path())
            assert doc.paragraphs
            if title == '修订稿':
                assert '验收合格后30日付款' in doc.paragraphs[0].text
            expect(menu).to_have_count(0)
        page.reload()
        page.locator('#card-payment').get_by_role('button', name='拒绝', exact=True).click()
        export.click()
        expect(page.get_by_text('请先完成审核，再导出文件', exact=True)).to_be_visible()
        assert page.request.get(base + '/api/tasks/' + task['id']).json()['status'] == 'pending_review'
        assert not errors, errors
        browser.close()


def test_export_before_review_shows_message(isolated_workspace):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with live_server() as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        task = page.request.post(base + '/api/tasks', data={'title': '未审查合同', 'text': '第一条 验收后付款。'}).json()
        page.goto(base + '/workbench?task=' + task['id'])
        page.get_by_role('button', name='导出', exact=True).click()
        expect(page.get_by_text('请先完成审核，再导出文件', exact=True)).to_be_visible()
        expect(page.get_by_role('menu')).to_have_count(0)
        browser.close()
