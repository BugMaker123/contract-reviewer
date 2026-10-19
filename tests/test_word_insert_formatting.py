from io import BytesIO
import pytest
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from backend import workspace_api, word_export
from backend.source_storage import save_source


@pytest.mark.parametrize('tracked', [False, True])
@pytest.mark.parametrize('paragraph_mark_only', [False, True])
def test_inserted_clause_inherits_body_font_spacing_and_indent(isolated_workspace, tracked, paragraph_mark_only):
    doc = Document()
    body = doc.add_paragraph('10.4 本合同自双方签章起生效。')
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Pt(24)
    properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for field in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn('w:' + field), '宋体')
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '24')
    properties.append(fonts); properties.append(size)
    if paragraph_mark_only:
        body._p.get_or_add_pPr().append(properties)
    else:
        body.runs[0]._r.insert(0, properties)
    signature = doc.add_paragraph('甲方：公司；乙方：公司。')
    signature.runs[0].font.size = Pt(9)
    text = '\n'.join(p.text for p in doc.paragraphs)
    buffer = BytesIO(); doc.save(buffer)
    source = save_source(workspace_api.repo, word_export.SOURCES, {'extension': '.docx', 'text': text}, buffer.getvalue())
    run = {'input': {'title': '合同', 'text': text, 'source_id': source['id']},
           'result': {'issues': []}, 'created_at': '2026-10-19T00:00:00+00:00', 'applied': [], 'comments': {}}
    inserted = '10.5 甲方保证所供产品不侵犯第三方知识产权。'
    final = body.text + '\n' + inserted + '\n' + signature.text
    result, layout = word_export.contract_document(run, final, workspace_api.repo, workspace_api.word_document, tracked=tracked)
    assert layout == 'preserved'
    paragraph = result.paragraphs[1]
    assert word_export.visible_paragraph_text(paragraph._p) == inserted
    assert paragraph.paragraph_format.line_spacing == 1.5
    assert paragraph.paragraph_format.first_line_indent.pt == 24
    rpr = paragraph._p.find('.//' + qn('w:rPr'))
    assert rpr.find(qn('w:sz')).get(qn('w:val')) == '24'
    assert rpr.find(qn('w:rFonts')).get(qn('w:eastAsia')) == '宋体'
    assert result.paragraphs[2].runs[0].font.size.pt == 9
    if tracked:
        assert paragraph._p.find(qn('w:ins')) is not None
