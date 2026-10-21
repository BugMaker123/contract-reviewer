from test_tracked_revisions import revision_client, seed_review
from test_ui_composition import live_server
from playwright.sync_api import sync_playwright, expect


def test_accept_then_reject_keeps_original_and_other_pending_markup(revision_client):
    task, run = seed_review(revision_client, True)
    with live_server() as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width':1920, 'height':1000})
        page.goto(base+'/workbench?task='+task['id'])
        preview = page.frame_locator('iframe[title="合同修订文档"]')
        card = page.locator('#card-payment')
        expect(preview.locator('[data-review-issue="payment"] ins')).not_to_have_count(0)
        card.get_by_role('button', name='采纳', exact=True).click()
        expect(preview.locator('[data-review-issue="payment"]')).to_have_class(__import__('re').compile('review-applied'))
        card.get_by_role('button', name='拒绝', exact=True).click()
        expect(preview.locator('[data-review-issue="payment"]')).to_have_class(__import__('re').compile('review-rejected'))
        expect(preview.locator('#review-callout')).to_have_text('已拒绝 · 保留原文')
        expect(preview.locator('[data-review-issue="payment"]')).to_contain_text('验收后付款')
        expect(preview.locator('[data-review-issue="payment"] ins')).to_have_count(0)
        expect(preview.locator('[data-review-issue="liability"] ins')).not_to_have_count(0)
        expect(preview.locator('[data-review-issue="liability"] del')).not_to_have_count(0)
        assert preview.locator('[data-review-issue="liability"] del').first.evaluate('(el)=>getComputedStyle(el).textDecorationLine') == 'line-through'
        page.reload()
        expect(preview.locator('[data-review-issue="payment"]')).to_have_class(__import__('re').compile('review-rejected'))
        expect(preview.locator('[data-review-issue="liability"] ins')).not_to_have_count(0)
        browser.close()


def test_blank_page_break_does_not_join_next_clause_selection(revision_client):
    from io import BytesIO
    from docx import Document
    from backend import workspace_api
    task, run = seed_review(revision_client)
    document = Document()
    document.add_paragraph(run['result']['issues'][0]['redline']['original_text'])
    empty = document.add_paragraph('')
    empty.paragraph_format.page_break_before = True
    document.add_paragraph(run['result']['issues'][1]['redline']['original_text'])
    output = BytesIO(); document.save(output)
    imported = revision_client.post('/api/contracts/import',files={'file':('paged.docx',output.getvalue())}).json()
    changes = {'text':imported['text'],'source_id':imported['source_id']}
    workspace_api.repo.update(task['id'],changes)
    workspace_api.repo.update(run['id'],{'input':{**run['input'],**changes}})
    with live_server() as base, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1920,'height':1200})
        page.goto(base+'/workbench?task='+task['id'])
        preview=page.frame_locator('iframe[title="合同修订文档"]')
        expect(preview.locator('[data-review-issue="liability"]')).not_to_have_count(0)
        assert preview.locator('[data-review-issue="liability"]').evaluate_all('(els)=>els.every(el=>el.textContent.trim().length>0)')
        page.locator('#card-liability .prototype-locate').click()
        expect(preview.locator('#review-selection')).to_be_visible()
        assert preview.locator('#review-selection').bounding_box()['height'] < 200
        browser.close()


def test_appended_paragraphs_belong_to_same_revision_selection(revision_client):
    from backend import workspace_api
    task, run = seed_review(revision_client)
    issue = run['result']['issues'][0]
    addition = '新增付款保障条款。'
    issue['redline']['original_text'] = run['input']['text']
    issue['redline']['suggested_text'] = run['input']['text'] + '\n' + addition
    run['result']['issues'] = [issue]
    workspace_api.repo.update(run['id'], {'result':run['result']})
    with live_server() as base, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1920,'height':1000})
        page.goto(base+'/workbench?task='+task['id'])
        preview=page.frame_locator('iframe[title="合同修订文档"]')
        expect(preview.locator('[data-review-issue="payment"]')).to_have_count(3)
        expect(preview.locator('[data-review-issue="payment"] ins')).to_have_text(addition)
        page.locator('#card-payment .prototype-locate').click()
        expect(preview.locator('.review-active ins')).to_have_text(addition)
        browser.close()


def test_missing_fields_are_risk_annotations_and_survive_acceptance(revision_client):
    from backend import workspace_api
    task, run = seed_review(revision_client)
    original = '法定代表人：\n地址：\n法定代表人：王彬\n地址：实际地址'
    issue=run['result']['issues'][0]
    issue.update({'issue_title':'法定代表人及地址空白', 'risk_analysis':'甲方法定代表人、地址缺失，需要补充。', 'evidence':original,
                  'redline':{'original_text':original,'suggested_text':original+'\n双方应核实主体信息。','change_type':'replace'}})
    workspace_api.repo.update(task['id'],{'text':original})
    workspace_api.repo.update(run['id'],{'input':{**run['input'],'text':original},'result':{**run['result'],'issues':[issue]}})
    with live_server() as base, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1920,'height':1000})
        page.goto(base+'/workbench?task='+task['id'])
        preview=page.frame_locator('iframe[title="合同修订文档"]')
        expect(preview.locator('.review-missing-field')).to_have_count(2)
        expect(preview.locator('.review-missing-field').first).to_have_attribute('data-missing-fields','法定代表人')
        assert preview.locator('.review-missing-field').first.evaluate('(el)=>getComputedStyle(el,"::after").content').find('字段缺失')>=0
        expect(preview.locator('.review-missing-field').filter(has_text='王彬')).to_have_count(0)
        page.locator('#card-payment').get_by_role('button',name='采纳',exact=True).click()
        expect(preview.locator('.review-applied')).not_to_have_count(0)
        expect(preview.locator('.review-missing-field')).to_have_count(2)
        assert '字段缺失' not in preview.locator('article').first.text_content()
        browser.close()
