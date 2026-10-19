"""Second-round workflow, version provenance and rich clipboard regressions."""
from io import BytesIO
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient
from backend.main import app
from backend import graph, workspace_api
from backend.store import now

TEXT = '第一条 设备交付。\n第二条 验收后30日付款。'
HTML = '<h2 style="text-align:center">设备采购合同</h2><p><strong style="font-size:18pt;color:#c00000">第一条 设备交付。</strong></p><p>第二条 验收后30日付款。</p><table><tr><th>项目</th><th>说明</th></tr><tr><td>设备</td><td>一套</td></tr></table>'

@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(graph,'get_llm_client',lambda config:None)
    with TestClient(app) as client: yield client


def task(client, title='版本验收', **kwargs):
    response = client.post('/api/tasks',json={'title':title,'text':TEXT,'mode':'quick',**kwargs})
    assert response.status_code == 200,response.text
    return response.json()


def completed(client, **kwargs):
    item = task(client,**kwargs)
    run = workspace_api.repo.start_run(item['id'])
    result = {'issues':[], 'segments':[], 'summary':{'verdict':'待人工复核','risk_level':'一般','compliance_score':80},'coverage':'基础规则'}
    run = workspace_api.repo.update(run['id'],{'status':'completed','completed_at':now(),'result':result})
    workspace_api.repo.update(item['id'],{'status':'pending_review','active_run':run['id']})
    return client.get('/api/tasks/'+item['id']).json(),run


def accept_payment(client, item, run):
    run = client.post('/api/runs/'+run['id']+'/issues',json={'revision':run['revision'],'title':'付款期限','analysis':'约定付款时限','evidence':'验收后30日付款。','suggestion':'验收后15日付款。'}).json()
    issue = run['result']['issues'][-1]
    response = client.put('/api/runs/'+run['id']+'/handling',json={'revision':run['revision'],'applied':[issue['id']]})
    assert response.status_code == 200,response.text
    return response.json()


def test_task_paging_filtering_and_dates(client):
    items = [task(client,title=f'分页合同{i:02}') for i in range(23)]
    first = client.get('/api/tasks?page=1&page_size=10&archived=false').json()
    second = client.get('/api/tasks?page=2&page_size=10&archived=false').json()
    third = client.get('/api/tasks?page=3&page_size=10&archived=false').json()
    assert first['total'] == 23 and [len(p['items']) for p in (first,second,third)] == [10,10,3]
    assert not ({i['id'] for i in first['items']} & {i['id'] for i in second['items']})
    assert 'text' not in first['items'][0]
    query = client.get('/api/tasks?q=合同22&page=3&page_size=10&status=draft&mode=quick').json()
    assert query['total'] == 1 and query['page'] == 1
    archived = client.post('/api/tasks/'+items[-1]['id']+'/archive',json={'archived':True,'revision':items[-1]['revision']})
    assert archived.status_code == 200
    assert client.get('/api/tasks?archived=true').json()['total'] == 1
    assert client.get('/api/tasks?from_date=2026-12-31&to_date=2026-01-01').status_code == 422
    assert client.get('/api/tasks?page=0').status_code == 422


def test_rich_clipboard_becomes_formatted_source_and_export(client):
    response = client.post('/api/contracts/paste',json={'title':'粘贴样式','html':HTML + '<script>EVIL</script><img src="https://invalid.example/x">'})
    assert response.status_code == 200,response.text
    data = response.json(); assert 'EVIL' not in data['text']
    item,run = completed(client,text=data['text'],source_id=data['source_id'])
    response = client.get('/api/runs/'+run['id']+'/export?scope=current')
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    bold = next(r for p in document.paragraphs for r in p.runs if '设备交付' in r.text)
    assert bold.bold and bold.font.size.pt == 18 and str(bold.font.color.rgb) == 'C00000'
    assert document.tables[0].cell(1,0).text == '设备'
    assert response.headers['X-Word-Layout'] == 'preserved'
    assert client.post('/api/contracts/paste',json={'html':'<script>empty</script>'}).status_code == 422


def test_accepted_baseline_creates_new_snapshot_and_freezes_history(client):
    pasted = client.post('/api/contracts/paste',json={'html':HTML}).json()
    item,run = completed(client,text=pasted['text'],source_id=pasted['source_id'])
    run = accept_payment(client,item,run)
    original = client.get('/api/runs/'+run['id']).json()
    item = client.get('/api/tasks/'+item['id']).json()
    response = client.post('/api/tasks/'+item['id']+'/review/stream',json={'baseline':'accepted','task_revision':item['revision'],'run_revision':run['revision']})
    assert response.status_code == 200,response.text
    current = client.get('/api/tasks/'+item['id']).json()
    new = current['run']
    assert new['baseline_kind'] == 'accepted' and new['parent_run_id'] == run['id']
    assert new['number'] == run['number'] + 1 and '15日付款' in new['input']['text']
    assert new['input']['source_id'] != original['input']['source_id']
    assert client.get('/api/runs/'+run['id']).json() == original
    assert client.put('/api/runs/'+run['id']+'/handling',json={'revision':run['revision'],'applied':[]}).status_code == 409
    assert client.post('/api/runs/'+run['id']+'/followup',json={'issue_id':run['result']['issues'][0]['id'],'question':'说明'}).status_code == 409
    diff = client.get('/api/tasks/'+item['id']+'/comparison',params={'before_run':run['id'],'after_run':new['id']}).json()
    assert diff['kind'] == 'versions' and any('15日' in c['revised'] for c in diff['changes'])
    other = task(client,title='另一任务')
    assert client.get('/api/tasks/'+other['id']+'/comparison',params={'after_run':new['id']}).status_code == 422
    exported = client.get('/api/runs/'+new['id']+'/export?scope=current')
    doc = Document(BytesIO(exported.content))
    assert any(r.bold and r.font.size.pt == 18 for p in doc.paragraphs for r in p.runs)


def test_baseline_rejects_stale_versions_and_task_baseline_excludes_acceptance(client):
    item,run = completed(client); run = accept_payment(client,item,run)
    item = client.get('/api/tasks/'+item['id']).json()
    assert client.post('/api/tasks/'+item['id']+'/review/stream',json={'baseline':'accepted'}).status_code == 422
    assert client.post('/api/tasks/'+item['id']+'/review/stream',json={'baseline':'accepted','task_revision':item['revision'],'run_revision':run['revision']-1}).status_code == 409
    response = client.post('/api/tasks/'+item['id']+'/review/stream',json={'baseline':'task','task_revision':item['revision'],'run_revision':run['revision']})
    assert response.status_code == 200
    new = client.get('/api/tasks/'+item['id']).json()['run']
    assert new['input']['text'] == TEXT and new['baseline_kind'] == 'task'


def test_browser_upload_and_rich_paste(isolated_workspace, monkeypatch):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    monkeypatch.setattr(graph,'get_llm_client',lambda config:None)
    with live_server() as base,sync_playwright() as p:
        browser = p.chromium.launch(); page = browser.new_page(viewport={'width':1600,'height':1000})
        errors = []; page.on('pageerror',lambda err:errors.append(str(err)))
        page.goto(base)
        expect(page.get_by_role('button',name='📑 示例合同',exact=True)).to_have_count(0)
        with page.expect_file_chooser() as chooser:
            page.get_by_role('button',name='📤 上传合同',exact=True).click()
        chooser.value.set_files({'name':'上传按钮.txt','mimeType':'text/plain','buffer':TEXT.encode()})
        expect(page.get_by_role('button',name='预览合同',exact=True)).to_be_visible()
        page.get_by_role('button',name='✍️ 粘贴正文',exact=True).click()
        editor = page.get_by_role('textbox',name='合同正文',exact=True)
        editor.evaluate("(el,html)=>{el.focus();const data=new DataTransfer();data.setData('text/html',html);el.dispatchEvent(new ClipboardEvent('paste',{clipboardData:data,bubbles:true,cancelable:true}));}",HTML)
        expect(editor.locator('strong')).to_contain_text('设备交付')
        expect(editor.locator('table')).to_have_count(1)
        page.get_by_placeholder('填写便于识别的合同名称').fill('富文本浏览器验收')
        page.get_by_role('button',name='进入工作台核对合同',exact=True).click()
        page.wait_for_url('**/workbench?task=*')
        task_id = page.url.split('task=')[1]
        item = page.request.get(base+'/api/tasks/'+task_id).json()
        assert item['source_id'] and '设备交付' in item['text']
        expect(page.frame_locator('iframe[title="合同原件预览"]').locator('article table')).to_be_visible()
        page.get_by_role('link',name='退出工作台',exact=True).click()
        expect(page.get_by_role('heading',name='新建合同审查',exact=True)).to_be_visible()
        expect(page.locator('.recent-work')).to_contain_text('富文本浏览器验收')
        page.locator('.recent-work').get_by_role('link',name='继续审查 →',exact=True).click()
        expect(page.get_by_role('heading',name='富文本浏览器验收',exact=True)).to_be_visible()
        assert not errors,errors
        browser.close()


def test_browser_records_paging_comparison_and_completion(isolated_workspace, monkeypatch):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    monkeypatch.setattr(graph,'get_llm_client',lambda config:None)
    with live_server() as base,sync_playwright() as p:
        browser = p.chromium.launch(); page = browser.new_page(viewport={'width':1600,'height':1000})
        errors = []; page.on('pageerror',lambda err:errors.append(str(err)))
        for i in range(12): page.request.post(base+'/api/tasks',data={'title':f'列表分页{i:02}','text':TEXT,'mode':'quick'})
        with TestClient(app) as client:
            item,run = completed(client,title='记录交付验收'); run = accept_payment(client,item,run)
        page.goto(base+'/records')
        expect(page.locator('tbody tr')).to_have_count(10)
        page.get_by_role('button',name='下一页',exact=True).click()
        expect(page.locator('tbody tr')).to_have_count(3)
        page.get_by_role('textbox',name='搜索审查记录',exact=True).fill('记录交付验收')
        expect(page.locator('tbody tr')).to_have_count(1)
        expect(page.get_by_role('button',name='上一页',exact=True)).to_be_disabled()
        row = page.locator('tbody tr').first
        with page.expect_download() as download: row.get_by_role('link',name='导出审核结果',exact=True).click()
        assert download.value.suggested_filename.endswith('.docx')
        row.get_by_role('button',name='前后对比',exact=True).click()
        expect(page.get_by_role('dialog',name='审核前后对比')).to_contain_text('验收后15日付款')
        page.get_by_role('button',name='关闭对比',exact=True).click()
        row.get_by_role('link',name='打开',exact=True).click()
        page.get_by_role('button',name='完成审核',exact=True).click()
        expect(page.get_by_role('heading',name='新建合同审查',exact=True)).to_be_visible()
        expect(page.locator('.recent-work')).to_contain_text('已完成')
        expect(page.locator('.recent-work').get_by_role('link',name='查看审核结果 →',exact=True)).to_be_visible()
        assert page.request.get(base+'/api/tasks/'+item['id']).json()['status'] == 'completed'
        assert not errors,errors
        browser.close()


def test_followup_valid_request_is_saved(client, monkeypatch):
    item,run = completed(client); run = accept_payment(client,item,run)
    async def answer(config,prompt): return {'answer':'请核对付款日期','suggested_text':''}
    monkeypatch.setattr(workspace_api,'model_json',answer)
    identifier = run['result']['issues'][0]['id']
    response = client.post('/api/runs/'+run['id']+'/followup',json={'issue_id':identifier,'question':'请说明付款约定'})
    assert response.status_code == 200,response.text
    assert response.json()['followups'][identifier][0]['answer'] == '请核对付款日期'


def test_browser_rereview_baseline_choice_and_history(isolated_workspace, monkeypatch):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    monkeypatch.setattr(graph,'get_llm_client',lambda config:None)
    with live_server() as base,sync_playwright() as p:
        with TestClient(app) as client:
            item,run = completed(client); run = accept_payment(client,item,run)
        browser = p.chromium.launch(); page = browser.new_page(viewport={'width':1920,'height':1080})
        errors = []; page.on('pageerror',lambda err:errors.append(str(err)))
        page.goto(base+'/workbench?task='+item['id'])
        page.get_by_role('button',name='重新审查',exact=True).click()
        dialog = page.get_by_role('dialog',name='选择重新审查稿件')
        expect(dialog.locator('input[value="accepted"]')).to_be_checked()
        dialog.get_by_role('button',name='确认并开始审查',exact=True).click()
        page.get_by_role('button',name='查看审查修订稿',exact=True).click()
        latest = page.request.get(base+'/api/tasks/'+item['id']).json()
        assert latest['run']['baseline_kind'] == 'accepted'
        assert '验收后15日付款' in latest['run']['input']['text']
        expect(page.locator('.wb-functional-overview')).to_contain_text('上一版已采纳修订稿')
        page.locator('.wb-functional-overview > summary').click()
        page.get_by_label('审查历史版本',exact=True).select_option(run['id'])
        expect(page.get_by_text('正在查看历史审查版本，原文和结果保留不变。',exact=False)).to_be_visible()
        expect(page.get_by_role('button',name='编辑正文',exact=True)).to_be_disabled()
        page.get_by_role('link',name='返回当前任务',exact=True).click()
        expect(page.get_by_role('button',name='编辑正文',exact=True)).to_be_enabled()
        assert not errors,errors
        page.screenshot(path='artifacts/version-baseline-workbench.png')
        browser.close()
