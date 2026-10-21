from io import BytesIO
import pytest
from docx import Document
from fastapi.testclient import TestClient
from backend.main import app
from backend import template_review


def upload(client, text):
    doc = Document()
    doc.add_paragraph(text)
    data = BytesIO()
    doc.save(data)
    return client.post('/api/contracts/import', files={'file': ('template.docx', data.getvalue())}).json()['source_id']


def setup_template(client):
    group = client.post('/api/template-groups', json={'name': '测试分组'}).json()
    source = upload(client, '甲方：____；期限：30日。')
    record = client.post('/api/contract-templates', json={'name': '模板', 'group_id': group['id'], 'source_id': source}).json()
    return record, '/api/contract-templates/' + record['id']


def region(**changes):
    return {'id': 'f1', 'name': '甲方', 'block_id': 'b0', 'start': 3, 'end': 7,
            'type': 'text', 'policy': 'conditional', 'confirmed': True, **changes}


def test_rules_publication_trial_and_replacement():
    with TestClient(app) as client:
        record, base = setup_template(client)
        detail = client.get(base+'/review').json()
        assert detail['blocks'][0]['text'] == '甲方：____；期限：30日。'
        assert client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': [region(), region(id='overlap', start=4)]}).status_code == 422
        saved = client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': [region()]}).json()
        assert saved['revision'] == 2
        assert client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': []}).status_code == 409
        published = client.post(base+'/publish', json={'revision': 2}).json()
        assert published['review_status'] == 'published'
        release = template_review.repo().get(published['release_id'], 'template_release')
        good = upload(client, '甲方：某某公司；期限：30日。')
        result = client.post(base+'/trial', json={'revision': 3, 'source_id': good}).json()
        assert result['result'] == 'PASS' and result['values'][0]['value'] == '某某公司'
        assert result['eligible_for_skip'] is False
        bad = upload(client, '甲方：某某公司；期限：60日。')
        assert client.post(base+'/trial', json={'revision': 3, 'source_id': bad}).json()['result'] == 'DIFFERENT'
        assert client.post(base+'/trial', json={'revision': 3, 'source_id': record['source_id']}).json()['result'] == 'DIFFERENT'
        client.put(base+'/review', json={'revision': 3, 'source_id': record['source_id'], 'fields': [region(confirmed=False)]})
        assert client.post(base+'/publish', json={'revision': 4}).status_code == 422
        assert template_review.repo().get(release['id'])['fields'][0]['confirmed'] is True
        updated = client.put(base, json={'name': '新模板', 'group_id': record['group_id'], 'source_id': good, 'revision': 4}).json()
        assert updated['review_status'] == 'draft'
        assert client.get(base+'/review').json()['fields'] == []
        assert len(client.get(base+'/review').json()['releases']) == 1


def test_ai_positions_validated_and_no_implicit_save(monkeypatch):
    async def fake_model(config, prompt):
        return {'fields': [{'name': '甲方', 'block_id': 'b0', 'text': '____', 'type': 'text'}]}
    monkeypatch.setattr(template_review, 'model_json', fake_model)
    with TestClient(app) as client:
        record, base = setup_template(client)
        result = client.post(base+'/extract-fields', json={'revision': 1}).json()
        assert result['fields'][0]['start'] == 3
        assert not result['fields'][0]['confirmed']
        assert result['fields'][0]['policy'] == 'manual'
        assert client.get(base+'/review').json()['fields'] == []
        async def invalid_model(config, prompt):
            return {'fields': [{'name': '编造', 'block_id': 'b0', 'text': '不存在'}]}
        monkeypatch.setattr(template_review, 'model_json', invalid_model)
        assert client.post(base+'/extract-fields', json={'revision': 1}).status_code == 422


def test_numeric_and_structure_matching():
    before = [{'id': 'b0', 'text': '天数：____天', 'section': '正文', 'kind': 'paragraph'}]
    fields = [template_review.Region(**region(type='integer', minimum=1, maximum=30)).model_dump()]
    assert template_review.compare_blocks(before, [{**before[0], 'text': '天数：20天'}], fields)['result'] == 'PASS'
    assert template_review.compare_blocks(before, [{**before[0], 'text': '天数：60天'}], fields)['result'] == 'DIFFERENT'
    assert template_review.compare_blocks(before, [], fields)['result'] == 'UNDETERMINED'
    fields[0]['policy'] = 'manual'
    assert template_review.compare_blocks(before, [{**before[0], 'text': '天数：20天'}], fields)['result'] == 'DIFFERENT'


def test_review_browser_manual_mark_persist_publish():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with TestClient(app) as client:
        record, base = setup_template(client)
    with live_server() as address, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1700, 'height': 1100})
        errors = []
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.goto(address+'/templates/'+record['id']+'/review')
        expect(page.get_by_role('heading', name='模板')).to_be_visible()
        expect(page.locator('.portal-sidebar')).to_have_count(0)
        assert page.locator('.tr-workspace').bounding_box()['y'] <= 150
        expect(page.get_by_role('link', name='返回主页', exact=True)).to_have_attribute('href', '/')
        page.get_by_role('button', name='手动添加', exact=True).click()
        dialog = page.get_by_role('dialog')
        dialog.get_by_label('原文位置', exact=True).select_option('b0')
        dialog.get_by_label('开始位置', exact=True).fill('4')
        dialog.get_by_label('结束位置（包含）', exact=True).fill('7')
        dialog.get_by_role('button', name='保存位置').click()
        page.get_by_label('字段名称', exact=True).fill('甲方名称')
        page.get_by_label('免审处理', exact=True).select_option('conditional')
        expect(page.get_by_role('button', name='验证合同', exact=True)).to_be_disabled()
        expect(page.locator('.tr-readiness')).to_have_text('待核验 1 项')
        expect(page.get_by_role('button', name='发布版本', exact=True)).to_have_attribute('title', '还有 1 个区域待核验，请在右侧逐项点击“确认并下一项”')
        page.get_by_role('button', name='确认并下一项').click()
        expect(page.get_by_role('button', name='验证合同', exact=True)).to_be_enabled()
        expect(page.locator('.tr-readiness')).to_have_text('全部已确认')
        page.get_by_role('button', name='保存草稿').click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        expect(page.get_by_role('status')).to_contain_text('草稿已保存，共 1 个填写区域')
        expect(page.locator('.tr-foot time')).to_contain_text('上次保存')
        page.reload()
        expect(page.locator('.tr-field')).to_contain_text('甲方名称')
        page.get_by_role('button', name='发布版本', exact=True).click()
        page.get_by_role('dialog').get_by_role('button', name='确认', exact=True).click()
        expect(page.locator('.tr-heading .tr-status-badge')).to_have_text('已发布')
        assert not errors
        browser.close()


def test_ai_partial_results_and_typographic_location(monkeypatch):
    async def fake_model(config, prompt):
        return {'fields': [
            {'name': '甲方', 'block_id': 'b0', 'text': '＿＿＿＿'},
            {'name': '重复', 'block_id': 'b0', 'text': '____'},
            {'name': '编造', 'block_id': 'b0', 'text': '不存在'},
            {'name': '坏类型', 'block_id': 'b0', 'text': '30', 'type': 'text/integer'},
        ]}
    monkeypatch.setattr(template_review, 'model_json', fake_model)
    with TestClient(app) as client:
        record, base = setup_template(client)
        response = client.post(base+'/extract-fields', json={'revision': 1})
        assert response.status_code == 200
        data = response.json()
        assert len(data['fields']) == 1
        assert data['fields'][0]['start'] == 3 and data['fields'][0]['end'] == 7
        assert data['fields'][0]['policy'] == 'manual' and not data['fields'][0]['confirmed']
        assert len(data['skipped']) == 2
        assert client.get(base+'/review').json()['fields'] == []


def test_ai_failure_distinguished_and_draft_preserved(monkeypatch):
    async def failed_model(config, prompt):
        raise RuntimeError('provider-private-key')
    monkeypatch.setattr(template_review, 'model_json', failed_model)
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': [region()]})
        response = client.post(base+'/extract-fields', json={'revision': 2})
        assert response.status_code == 502
        assert '模型调用失败' in response.json()['detail']
        assert 'provider-private-key' not in response.text
        assert len(client.get(base+'/review').json()['fields']) == 1


def test_ai_ambiguous_normalization_not_guessed(monkeypatch):
    async def fake_model(config, prompt):
        return {'fields': [{'name': '甲方', 'block_id': 'b0', 'text': '＿＿'}]}
    monkeypatch.setattr(template_review, 'model_json', fake_model)
    with TestClient(app) as client:
        record, base = setup_template(client)
        response = client.post(base+'/extract-fields', json={'revision': 1})
        assert response.status_code == 422
        assert '无法唯一定位' in response.json()['detail']


def test_review_browser_extraction_feedback(monkeypatch):
    import asyncio
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect

    async def fake_model(config, prompt):
        await asyncio.sleep(0.6)
        return {'fields': [
            {'name': '甲方名称', 'block_id': 'b0', 'text': '____'},
            {'name': '无法定位字段', 'block_id': 'b0', 'text': '不在原文中'},
        ]}
    monkeypatch.setattr(template_review, 'model_json', fake_model)
    with TestClient(app) as client:
        record, base = setup_template(client)
    with live_server() as address, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1700, 'height': 1100})
        page.goto(address+'/templates/'+record['id']+'/review')
        page.get_by_role('button', name='AI 提取', exact=True).click()
        page.get_by_role('dialog').get_by_role('button', name='确认', exact=True).click()
        expect(page.locator('.tr-extraction-status')).to_contain_text('正在识别合同填写区域')
        expect(page.locator('.tr-field')).to_contain_text('甲方名称')
        expect(page.locator('.tr-extraction-status')).to_contain_text('已识别 1 个填写区域')
        page.locator('.tr-extraction-status summary').click()
        expect(page.locator('.tr-extraction-status')).to_contain_text('无法定位字段')
        expect(page.locator('.tr-mark')).to_have_text('{{ 甲方名称 }}')
        page.get_by_label('搜索填写区域').fill('找不到')
        expect(page.locator('.tr-no-matches')).to_have_text('没有匹配的填写区域')
        page.get_by_label('搜索填写区域').fill('')
        expect(page.locator('.tr-field')).to_have_count(1)
        page.locator('.tr-view-options > summary').click()
        page.get_by_label('显示段落位置').check()
        expect(page.locator('.tr-block-label').first).to_be_visible()
        page.get_by_label('显示段落位置').uncheck()
        expect(page.locator('.tr-block-label')).to_have_count(0)
        expect(page.locator('.tr-foot span')).to_have_text('尚未保存')
        browser.close()


def test_review_browser_save_failure_preserves_edits():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': [region()]})
    with live_server() as address, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1536, 'height': 864})
        page.goto(address+'/templates/'+record['id']+'/review')
        page.get_by_label('字段名称', exact=True).fill('修改后字段')
        def conflict(request_route):
            if request_route.request.method == 'PUT':
                request_route.fulfill(status=409, content_type='application/json', body='{"detail":"模板已被修改，请刷新后重试"}')
            else:
                request_route.continue_()
        page.route('**/api/contract-templates/*/review', conflict)
        page.get_by_role('button', name='保存草稿', exact=True).click()
        expect(page.get_by_role('alert')).to_contain_text('保存失败：模板已被修改')
        expect(page.get_by_label('字段名称', exact=True)).to_have_value('修改后字段')
        expect(page.locator('.tr-foot span')).to_have_text('尚未保存')
        page.unroute('**/api/contract-templates/*/review', conflict)
        page.get_by_role('button', name='保存草稿', exact=True).click()
        expect(page.get_by_role('status')).to_contain_text('草稿已保存')
        page.reload()
        expect(page.get_by_label('字段名称', exact=True)).to_have_value('修改后字段')
        browser.close()


def test_browser_bulk_confirm_and_cancel_persist():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with TestClient(app) as client:
        record, base = setup_template(client)
        fields = [region(confirmed=False), region(id='f2', name='期限', start=11, end=13, confirmed=False)]
        response = client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': fields})
        assert response.status_code == 200
    with live_server() as address, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width':1536,'height':864})
        page.goto(address+'/templates/'+record['id']+'/review')
        expect(page.locator('.tr-fields-panel .tr-bulk-actions')).to_have_count(0)
        expect(page.locator('.tr-inspector').get_by_role('button', name='全部确认', exact=True)).to_be_visible()
        expect(page.get_by_role('button', name='取消全部确认', exact=True)).to_be_disabled()
        page.get_by_label('搜索填写区域').fill('甲方')
        expect(page.locator('.tr-field')).to_have_count(1)
        page.get_by_role('button', name='全部确认', exact=True).click()
        expect(page.locator('.tr-readiness')).to_have_text('全部已确认')
        expect(page.get_by_role('button', name='验证合同', exact=True)).to_be_enabled()
        expect(page.get_by_role('button', name='发布版本', exact=True)).to_be_enabled()
        page.get_by_role('button', name='保存草稿', exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        page.reload()
        expect(page.locator('.tr-readiness')).to_have_text('全部已确认')
        expect(page.get_by_role('button', name='全部确认', exact=True)).to_be_disabled()
        page.get_by_role('button', name='取消全部确认', exact=True).click()
        expect(page.locator('.tr-readiness')).to_have_text('待核验 2 项')
        expect(page.get_by_role('button', name='发布版本', exact=True)).to_be_disabled()
        page.get_by_role('button', name='保存草稿', exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        page.reload()
        expect(page.locator('.tr-readiness')).to_have_text('待核验 2 项')
        page.get_by_role('button', name='全部确认', exact=True).click()
        page.get_by_label('字段名称', exact=True).fill('改动后的甲方')
        expect(page.locator('.tr-readiness')).to_have_text('待核验 1 项')
        browser.close()


@pytest.mark.parametrize('policy,text,status,manual,issues', [
    ('conditional', '甲方：某公司；期限：30日。', 'pass', 0, 0),
    ('manual', '甲方：某公司；期限：30日。', 'needs_review', 1, 0),
    ('conditional', '甲方：____；期限：30日。', 'needs_changes', 0, 1),
])
def test_validation_report_explains_field_results(policy, text, status, manual, issues):
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review', json={'revision':1,'source_id':record['source_id'],'fields':[region(policy=policy)]})
        candidate = upload(client, text)
        result = client.post(base+'/trial', json={'revision':2,'source_id':candidate}).json()
        report = result['report']
        assert report['status'] == status
        assert report['counts']['manual'] == manual
        assert report['counts']['issue'] == issues
        assert report['counts']['total'] == 1
        assert report['fields'][0]['template_value'] == '____'
        assert report['fields'][0]['unchanged'] is (issues == 1)
        assert result['eligible_for_skip'] is False


def test_validation_report_preserves_issues_with_complex_document(monkeypatch):
    original = template_review.document_blocks
    def with_warning(source):
        result = original(source)
        result['warnings'] = ['合同含图片，图片内容未检查']
        return result
    monkeypatch.setattr(template_review, 'document_blocks', with_warning)
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review',json={'revision':1,'source_id':record['source_id'],'fields':[region()]})
        candidate=upload(client,'甲方：____；期限：30日。')
        report=client.post(base+'/trial',json={'revision':2,'source_id':candidate}).json()['report']
        assert report['status']=='needs_changes'
        assert report['counts']['issue']==1
        assert report['counts']['differences']==0
        assert len(report['limitations'])==1
        candidate=upload(client,'完全不同的段落')
        report=client.post(base+'/trial',json={'revision':2,'source_id':candidate}).json()['report']
        assert report['counts']['not_checked']==1
        assert report['fields'][0]['value'] is None


def test_browser_validation_report_and_return_to_region(monkeypatch):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    original = template_review.document_blocks
    def with_warning(source):
        result=original(source)
        result['warnings']=['合同含图片，图片内容未检查']
        return result
    monkeypatch.setattr(template_review,'document_blocks',with_warning)
    with TestClient(app) as client:
        record,base=setup_template(client)
        client.put(base+'/review',json={'revision':1,'source_id':record['source_id'],'fields':[region(policy='manual')]})
    doc=Document();doc.add_paragraph('甲方：某某公司；期限：30日。')
    data=BytesIO();doc.save(data)
    with live_server() as address,sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1536,'height':864})
        page.goto(address+'/templates/'+record['id']+'/review')
        page.locator('input[type=file]').set_input_files({'name':'已填写的采购合同.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':data.getvalue()})
        dialog=page.get_by_role('dialog',name='合同验证结果')
        expect(dialog).to_be_visible()
        expect(dialog.get_by_role('heading',name='检查未完成，部分内容需人工核对')).to_be_visible()
        expect(dialog).to_contain_text('已填写的采购合同.docx')
        expect(dialog.locator('.validation-table')).to_contain_text('某某公司')
        expect(dialog.locator('.validation-badge')).to_have_text('需人工审核')
        dialog.get_by_role('button',name='检查范围',exact=True).click()
        expect(dialog.locator('.validation-scope')).to_contain_text('图片内容未检查')
        dialog.get_by_role('button',name='全部填写内容').click()
        dialog.get_by_role('button',name='查看标记').click()
        expect(dialog).to_have_count(0)
        expect(page.locator('.tr-mark.selected')).to_be_visible()
        browser.close()


def test_browser_template_placeholders_preserve_source_offsets():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with TestClient(app) as client:
        record,base=setup_template(client)
        client.put(base+'/review',json={'revision':1,'source_id':record['source_id'],'fields':[region(name='甲方公司名称',confirmed=False)]})
    with live_server() as address,sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page(viewport={'width':1536,'height':864})
        page.goto(address+'/templates/'+record['id']+'/review')
        expect(page.locator('.tr-mark')).to_have_text('{{ 甲方公司名称 }}')
        expect(page.locator('.tr-rule .tr-quote')).to_have_count(0)
        page.get_by_label('字段名称',exact=True).fill('供方')
        expect(page.locator('.tr-mark')).to_have_text('{{ 供方 }}')
        page.locator('[data-block="b0"] [data-source-start="7"]').evaluate("""el=>{
            const range=document.createRange();range.setStart(el.firstChild,4);range.setEnd(el.firstChild,6);
            const selection=window.getSelection();selection.removeAllRanges();selection.addRange(range);
            el.dispatchEvent(new MouseEvent('mouseup',{bubbles:true}));
        }""")
        page.get_by_role('button',name='＋ 添加填写区域',exact=True).click()
        expect(page.locator('.tr-mark')).to_have_count(2)
        page.get_by_role('button',name='保存草稿',exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        with TestClient(app) as client:
            fields=client.get(base+'/review').json()['fields']
            assert fields[0]['start']==3 and fields[0]['end']==7
            assert fields[1]['start']==11 and fields[1]['end']==13
        page.get_by_role('button',name='原文对照',exact=True).click()
        expect(page.get_by_title('上传原文对照',exact=True)).to_be_visible()
        page.get_by_role('button',name='模板预览',exact=True).click()
        expect(page.locator('.tr-mark')).to_have_count(2)
        browser.close()


def test_blank_form_labels_and_signature_columns_are_covered():
    texts=['甲方：甲公司','法定代表人：','地址:', '乙方：乙公司','法定代表人：王某','地址：某市',
           '甲方：甲公司  乙方：乙公司','法定代表人或授权代表：             法定代表人或授权代表：','日期：                             日期：',
           '按照如下日期：收到货物后30日内付款。', '双方应在15日内提供书面证明。']
    blocks=[{'id':f'b{i}','text':t,'section':'正文','kind':'paragraph'} for i,t in enumerate(texts)]
    candidates=template_review.blank_field_candidates(blocks)
    fields=[f for f,_,_ in candidates]
    assert [f['name'] for f in fields]==['甲方法定代表人','甲方地址','甲方法定代表人或授权代表','乙方法定代表人或授权代表','甲方日期','乙方日期']
    assert fields[0]['start']==fields[0]['end']==6
    assert fields[1]['start']==fields[1]['end']==3
    assert all(not f['confirmed'] and f['policy']=='manual' for f in fields)
    template_review.validate_regions(fields,blocks)


def test_extraction_supplements_blanks_preserving_existing_rules(monkeypatch):
    async def fake_model(config,prompt):
        return {'fields':[{'name':'重命名甲方','block_id':'b0','text':'甲公司'},
                          {'name':'错误空栏','block_id':'b1','text':'法定代表人：'}]}
    monkeypatch.setattr(template_review,'model_json',fake_model)
    with TestClient(app) as client:
        record,base=setup_template(client)
        doc=Document()
        for text in ['甲方：甲公司','法定代表人：','地址:']:
            doc.add_paragraph(text)
        data=BytesIO();doc.save(data)
        source=client.post('/api/contracts/import',files={'file':('blank.docx',data.getvalue())}).json()['source_id']
        client.put(base,json={'name':'模板','group_id':record['group_id'],'source_id':source,'revision':1})
        existing=region(name='自定义甲方名称',start=3,end=6)
        client.put(base+'/review',json={'revision':2,'source_id':source,'fields':[existing]})
        response=client.post(base+'/extract-fields',json={'revision':3})
        assert response.status_code==200
        result=response.json()
        assert result['preserved']==1 and result['supplemented']==2
        assert len(result['fields'])==3
        assert result['fields'][0]['id']=='f1'
        assert result['fields'][0]['name']=='自定义甲方名称'
        assert result['fields'][0]['confirmed'] is True
        assert all(not f['confirmed'] for f in result['fields'][1:])
        assert template_review.rules_for(template_review.repo().get(record['id']))==[template_review.Region(**existing).model_dump()]


def test_empty_insertion_fields_can_save_publish_and_validate():
    with TestClient(app) as client:
        record,base=setup_template(client)
        source=upload(client,'法定代表人：')
        client.put(base,json={'name':'模板','group_id':record['group_id'],'source_id':source,'revision':1})
        field=region(name='法定代表人',start=6,end=6)
        response=client.put(base+'/review',json={'revision':2,'source_id':source,'fields':[field]})
        assert response.status_code==200
        filled=upload(client,'法定代表人：张三')
        result=client.post(base+'/trial',json={'revision':3,'source_id':filled}).json()
        assert result['result']=='PASS'
        assert result['values'][0]['value']=='张三'
        empty=client.post(base+'/trial',json={'revision':3,'source_id':source}).json()
        assert empty['report']['counts']['issue']==1
        duplicate=client.put(base+'/review',json={'revision':3,'source_id':source,'fields':[field,{**field,'id':'duplicate'}]})
        assert duplicate.status_code==422
        assert client.post(base+'/publish',json={'revision':3}).status_code==200


def test_browser_extract_empty_labels_and_manual_blank_insertion(monkeypatch):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright,expect
    async def fake_model(config,prompt):return {'fields':[]}
    monkeypatch.setattr(template_review,'model_json',fake_model)
    with TestClient(app) as client:
        record,base=setup_template(client)
        source=upload(client,'法定代表人：')
        client.put(base,json={'name':'模板','group_id':record['group_id'],'source_id':source,'revision':1})
    with live_server() as address,sync_playwright() as p:
        browser=p.chromium.launch();page=browser.new_page(viewport={'width':1536,'height':864})
        page.goto(address+'/templates/'+record['id']+'/review')
        page.get_by_role('button',name='AI 提取',exact=True).click()
        page.get_by_role('dialog').get_by_role('button',name='确认',exact=True).click()
        expect(page.locator('.tr-mark')).to_have_text('{{ 法定代表人 }}')
        page.get_by_role('button',name='全部确认',exact=True).click()
        page.get_by_role('button',name='保存草稿',exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        page.reload()
        expect(page.locator('.tr-mark')).to_have_text('{{ 法定代表人 }}')
        page.get_by_role('button',name='手动添加',exact=True).click()
        dialog=page.get_by_role('dialog')
        dialog.get_by_label('原文位置',exact=True).select_option('b0')
        dialog.get_by_label('在空栏插入字段（保留原文字）',exact=True).check()
        dialog.get_by_label('插入位置',exact=True).fill('0')
        dialog.get_by_role('button',name='保存位置',exact=True).click()
        expect(page.locator('.tr-mark')).to_have_count(2)
        page.get_by_role('button',name='保存草稿',exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        browser.close()


@pytest.mark.parametrize('text,label',[
    ('法定代表人：','法定代表人：'),
    ('地址:','地址:'),
    ('协议编号：','协议编号：'),
    ('甲方（供方）：山东某公司','甲方（供方）：'),
    ('法定代表人或授权代表：','法定代表人或授权代表：'),
])
def test_field_ranges_never_replace_fixed_labels(text,label):
    blocks=[{'id':'b0','text':text,'section':'正文','kind':'paragraph'}]
    original=template_review.Region(**region(start=0,end=len(text))).model_dump()
    fields,repaired=template_review.preserve_form_labels([original],blocks)
    assert repaired==1
    assert fields[0]['start']==len(label)
    assert fields[0]['end']==len(text)
    assert fields[0]['confirmed'] is False
    assert fields[0]['id']==original['id']
    template_review.validate_regions(fields,blocks)
    assert blocks[0]['text']==text


def test_browser_legacy_label_ranges_show_original_words():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright,expect
    with TestClient(app) as client:
        record,base=setup_template(client)
        source=upload(client,'法定代表人：')
        client.put(base,json={'name':'模板','group_id':record['group_id'],'source_id':source,'revision':1})
        document=template_review.document_blocks(source)
        template_review.repo().update(record['id'],{'review_draft':{'source_id':source,'sha256':document['sha256'],'fields':[template_review.Region(**region(name='甲方法定代表人',start=0,end=6)).model_dump()]}},2)
        detail=client.get(base+'/review').json()
        assert detail['repaired']==1
        assert detail['blocks'][0]['text']=='法定代表人：'
    with live_server() as address,sync_playwright() as p:
        browser=p.chromium.launch();page=browser.new_page(viewport={'width':1536,'height':864})
        page.goto(address+'/templates/'+record['id']+'/review')
        expect(page.locator('[data-block="b0"]')).to_have_text('法定代表人：{{ 甲方法定代表人 }}')
        expect(page.locator('.tr-foot span')).to_have_text('尚未保存')
        page.get_by_role('button',name='保存草稿',exact=True).click()
        expect(page.locator('.tr-foot span')).to_have_text('已保存')
        page.reload()
        expect(page.locator('[data-block="b0"]')).to_have_text('法定代表人：{{ 甲方法定代表人 }}')
        with TestClient(app) as client:
            saved=client.get(base+'/review').json()
            assert saved['fields'][0]['start']==6
            assert saved['fields'][0]['end']==6
            assert saved['blocks'][0]['text']=='法定代表人：'
        browser.close()


def test_open_existing_template_supplements_blanks_without_overwriting_draft():
    with TestClient(app) as client:
        record,base=setup_template(client)
        doc=Document()
        for text in ['甲方：甲公司','法定代表人：','地址:']:
            doc.add_paragraph(text)
        data=BytesIO();doc.save(data)
        source=client.post('/api/contracts/import',files={'file':('blank.docx',data.getvalue())}).json()['source_id']
        client.put(base,json={'name':'模板','group_id':record['group_id'],'source_id':source,'revision':1})
        existing=region(name='甲方公司名称',start=3,end=6)
        client.put(base+'/review',json={'revision':2,'source_id':source,'fields':[existing]})
        detail=client.get(base+'/review').json()
        assert detail['supplemented']==2
        assert len(detail['fields'])==3
        assert detail['fields'][0]['confirmed'] is True
        assert [f['name'] for f in detail['fields'][1:]]==['甲方法定代表人','甲方地址']
        assert all(f['origin']=='rule' and not f['confirmed'] for f in detail['fields'][1:])
        assert len(template_review.rules_for(template_review.repo().get(record['id'])))==1
        assert template_review.repo().get(record['id'])['revision']==3


def test_published_validation_uses_frozen_rules():
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review', json={'revision': 1, 'source_id': record['source_id'], 'fields': [region()]})
        published = client.post(base+'/publish', json={'revision': 2}).json()
        source = upload(client, '甲方：某某公司；期限：30日。')
        payload = {'source_id': source, 'release_id': published['release_id']}
        response = client.post(base+'/validate', json=payload)
        assert response.status_code == 200
        assert response.json()['report']['status'] == 'pass'
        assert response.json()['version'] == 1
        assert response.json()['passed'] is True
        # Draft edits make the template unavailable until it is published again.
        client.put(base+'/review', json={'revision': 3, 'source_id': record['source_id'], 'fields': [region(max_length=1)]})
        response = client.post(base+'/validate', json=payload)
        assert response.status_code == 409
        release = template_review.repo().get(published['release_id'], 'template_release')
        assert release['fields'][0]['max_length'] == 200
        updated = client.post(base+'/publish', json={'revision': 4}).json()
        assert client.post(base+'/validate', json=payload).status_code == 409
        payload['release_id'] = updated['release_id']
        assert client.post(base+'/validate', json=payload).json()['report']['status'] == 'needs_changes'


def test_published_validation_rejects_unpublished_and_wrong_release():
    with TestClient(app) as client:
        record, base = setup_template(client)
        assert client.post(base+'/validate', json={'source_id':record['source_id'], 'release_id':'missing'}).status_code == 409


def test_browser_review_entry_published_validation(tmp_path):
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with TestClient(app) as client:
        record, base = setup_template(client)
        client.put(base+'/review', json={'revision':1, 'source_id':record['source_id'], 'fields':[region()]})
        client.post(base+'/publish', json={'revision':2})
    path = tmp_path / 'entry-contract.docx'
    doc = Document()
    doc.add_paragraph('甲方：某某公司；期限：30日。')
    doc.save(path)
    with live_server() as address, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width':1440, 'height':1000})
        errors=[]
        imports=[]
        page.on('request', lambda r: imports.append(r.url) if r.url.endswith('/api/contracts/import') else None)
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.goto(address)
        expect(page.get_by_role('heading', name='选择本次审查方式')).to_be_visible()
        page.locator('.review-entry').evaluate('(el)=>el.getAnimations({subtree:true}).forEach(a=>a.finish())')
        page.screenshot(path='artifacts/prototypes/review-entry-implemented.png', full_page=True)
        page.get_by_role('button', name='已有标准模板核验').click()
        expect(page.locator('#baseline')).to_have_value(record['id'])
        page.locator('#validation-file').set_input_files(str(path))
        expect(page.get_by_role('button', name='开始模板合规核验')).to_be_enabled()
        page.get_by_role('button', name='开始模板合规核验').click()
        expect(page.get_by_role('heading', name='TRUE · 校验通过')).to_be_visible()
        expect(page.locator('.report-table-wrap tbody tr')).to_have_count(0)
        expect(page.get_by_text('没有字段未通过项', exact=True)).to_be_visible()
        page.screenshot(path='artifacts/prototypes/review-result-implemented.png', full_page=True)
        assert len(imports) == 1
        assert not errors
        browser.close()


def test_browser_review_entry_empty_templates():
    from test_ui_composition import live_server
    from playwright.sync_api import sync_playwright, expect
    with live_server() as address, sync_playwright() as p:
        browser=p.chromium.launch()
        page=browser.new_page()
        page.goto(address)
        page.get_by_role('button', name='已有标准模板核验').click()
        expect(page.get_by_text('暂无已发布模板。', exact=False)).to_be_visible()
        expect(page.get_by_role('button', name='开始模板合规核验')).to_have_count(0)
        expect(page.get_by_role('link', name='合同模板', exact=True).last).to_have_attribute('href','/templates')
        browser.close()
