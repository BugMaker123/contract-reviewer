from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend.main import app
from backend.sample_contracts import SAMPLE_CONTRACTS

client = TestClient(app)


def test_import_preserves_real_text_and_docx_table_order():
    content = '第一条 合作内容\n实际上传的短合同'
    response = client.post('/api/contracts/import', files={'file': ('合同.txt', content.encode('utf-8'))})
    assert response.status_code == 200
    assert response.json()['text'] == content
    document = Document()
    document.add_paragraph('第一条 实际合同')
    document.add_table(rows=1, cols=1).cell(0, 0).text = '表格付款条件'
    document.add_paragraph('第二条 交付')
    stream = BytesIO()
    document.save(stream)
    response = client.post('/api/contracts/import', files={'file': ('合同.docx', stream.getvalue())})
    assert response.json()['text'] == '第一条 实际合同\n表格付款条件\n第二条 交付'


def test_invalid_import_does_not_substitute_sample():
    for name, content, status in [('旧版.doc', b'invalid', 415), ('损坏.docx', b'broken', 422), ('空白.txt', b'  ', 422)]:
        response = client.post('/api/contracts/import', files={'file': (name, content)})
        assert response.status_code == status
        assert 'text' not in response.json()
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    stream = BytesIO()
    pdf.write(stream)
    response = client.post('/api/contracts/import', files={'file': ('扫描.pdf', stream.getvalue())})
    assert response.status_code == 422
    assert 'OCR' in response.json()['detail']


def test_stream_returns_all_nodes_and_complete_result(monkeypatch):
    monkeypatch.setattr('backend.graph.get_llm_client', lambda config: None)
    sample = next(iter(SAMPLE_CONTRACTS.values()))
    payload = {'contract_text': sample['text'], 'contract_title': sample['title']}
    sync = client.post('/api/review', json=payload).json()['data']
    response = client.post('/api/review/stream', json=payload)
    assert response.status_code == 200
    blocks = response.text.strip().split('\n\n')
    updates = [block for block in blocks if block.startswith('event: node_update')]
    assert len(updates) == 5
    complete = next(block for block in blocks if block.startswith('event: complete'))
    data = json.loads(complete.split('data: ', 1)[1])['data']
    assert data == sync
    assert data['tax_settlement']
    assert data['summary']


def test_rejects_empty_or_excessive_contract():
    assert client.post('/api/review', json={'contract_text': ''}).status_code == 422
    assert client.post('/api/review', json={'contract_text': 'a' * 200001}).status_code == 422


def test_import_supports_windows_unicode_text_and_rejects_binary_data():
    text = '第一条 标的\n第二条 付款条件'
    response = client.post('/api/contracts/import', files={'file': ('合同.txt', text.encode('utf-16'))})
    assert response.status_code == 200
    assert response.json()['text'] == text
    response = client.post('/api/contracts/import', files={'file': ('伪文本.txt', b'\x00binary\x00')})
    assert response.status_code == 422
