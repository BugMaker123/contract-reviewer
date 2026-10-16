"""Browser regression for edits, accepted revisions, and restored previews."""
import socket
import re
import threading

import uvicorn
from playwright.sync_api import sync_playwright, expect
from backend.main import app
from backend import workspace_api, graph
from backend.sample_contracts import SAMPLE_CONTRACTS


def open_opinions(page):
    expect(page.locator('#opinion-drawer')).to_be_visible()


def close_opinions(page):
    # Opinion handling is now a permanent workspace region.
    expect(page.locator('#opinion-drawer')).to_be_visible()


def test_preview_follows_edit_and_handling(isolated_workspace, monkeypatch):
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: None)
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.add_init_script("sessionStorage.setItem('LEXIS_REVIEW_DISPLAY','compare')")
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            base = f'http://127.0.0.1:{port}'
            page.goto(base)
            text = SAMPLE_CONTRACTS['business_cooperation_contract']['text']
            imported = page.request.post(base+'/api/contracts/import', multipart={'file': {'name':'test.txt', 'mimeType':'text/plain', 'buffer':text.encode('utf-8')}}).json()
            source = {'id': imported['source_id']}
            response = page.request.post(base + '/api/tasks', data={'title': '预览回归', 'type': '采购合同', 'text': text, 'mode': 'quick', 'source_id': source['id']})
            task = response.json()
            page.request.post(base + '/api/tasks/' + task['id'] + '/review/stream', data={})
            page.goto(base + '/workbench?task=' + task['id'])
            expect(page.locator('iframe[title="合同原件预览"]')).to_be_visible()
            page.get_by_role('button', name='重新审查', exact=True).click()
            page.get_by_role('button', name='查看审查修订稿', exact=True).click()
            expect(page.get_by_test_id('original-document-pane')).to_be_visible()
            expect(page.frame_locator('iframe[title="合同原件预览"]').locator('article')).not_to_have_count(0, timeout=30000)
            open_opinions(page)
            page.get_by_role('button', name='采纳').first.click()
            expect(page.locator('iframe[title="合同原件预览"]')).to_be_visible()
            expect(page.get_by_test_id('original-document-pane')).to_be_visible()
            preview = page.frame_locator('iframe[title="合同原件预览"]')
            expect(preview.locator('article')).not_to_have_count(0)
            saved_run = page.request.get(base+'/api/tasks/'+task['id']).json()['run']
            other = next(issue for issue in saved_run['result']['issues'] if issue['id'] not in saved_run['applied'])
            preview.locator('body').evaluate("() => { window.previewInstance = 'unchanged'; }")
            with page.expect_response(lambda response: '/handling' in response.url and response.request.method == 'PUT'):
                page.locator('#card-'+other['id']).get_by_role('button', name='不采纳', exact=True).click()
            expect(page.locator('#card-'+other['id']).get_by_text('处理状态：不采纳', exact=True)).to_be_visible()
            assert preview.locator('body').evaluate('() => window.previewInstance') == 'unchanged'

            page.reload()
            open_opinions(page)
            expect(page.get_by_test_id('original-document-pane')).to_be_visible()
            page.locator('.revision-more-actions>summary').first.click()
            page.get_by_role('button', name='撤销采纳').click()
            expect(page.get_by_role('button', name='撤销采纳', exact=True)).to_have_count(0)
            page.get_by_role('button', name='编辑条款', exact=True).first.click()
            editor = page.locator('.fixed textarea').last
            original = editor.input_value()
            updated = original + '【浏览器核验修改】'
            editor.fill(updated)
            page.get_by_role('button', name='保存条款', exact=True).click()
            expect(page.locator('#contract-viewer')).to_contain_text(updated)
            expect(page.locator('iframe[title="合同原件预览"]')).to_have_count(0)
            expect(page.locator('#contract-viewer .clause-active')).to_have_count(0)
            close_opinions(page)
            page.get_by_role('button', name='保存', exact=True).click()
            open_opinions(page)
            expect(page.get_by_text('工作内容已保存', exact=True)).to_be_visible()
            page.reload()
            open_opinions(page)
            expect(page.locator('#contract-viewer')).to_contain_text(updated)
            close_opinions(page)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            body_editor = page.get_by_role('textbox', name='编辑合同正文')
            body_editor.fill(body_editor.input_value() + '\n【整篇正文修改】')
            page.get_by_role('button', name='保存正文', exact=True).click()
            expect(page.locator('#contract-viewer')).to_contain_text('【整篇正文修改】')
            close_opinions(page)
            page.get_by_role('button', name='保存', exact=True).click()
            open_opinions(page)
            expect(page.get_by_text('工作内容已保存', exact=True)).to_be_visible()
            page.reload()
            open_opinions(page)
            expect(page.locator('#contract-viewer')).to_contain_text('【整篇正文修改】')
            page.get_by_role('button', name='重新审查', exact=True).click()
            close_opinions(page)
            expect(page.get_by_role('button', name='重新审查', exact=True)).to_be_enabled(timeout=30000)
            assert len(page.request.get(base+'/api/tasks/'+task['id']).json()['runs']) == 3
            page.reload()
            open_opinions(page)
            expect(page.locator('#contract-viewer')).to_contain_text('【整篇正文修改】')
            assert not errors, errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        sock.close()


def test_docx_inline_edit_stays_in_document(isolated_workspace, monkeypatch):
    monkeypatch.setattr(graph, "get_llm_client", lambda config: None)
    from io import BytesIO
    from docx import Document
    document = Document()
    p = document.add_paragraph()
    p.add_run('Payment: ').bold = True
    p.add_run('100 USD').italic = True
    for index in range(35):
        document.add_paragraph(f'Contract paragraph {index}: delivery and payment terms.')
    document.add_paragraph('Special paragraph\nwith a soft break')
    document.add_paragraph('Tabbed\tparagraph')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Delivery terms'
    data = BytesIO()
    document.save(data)
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={'width':1600,'height':1000})
            page.add_init_script("sessionStorage.setItem('LEXIS_REVIEW_DISPLAY','compare')")
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            base = f'http://127.0.0.1:{sock.getsockname()[1]}'
            imported = page.request.post(base+'/api/contracts/import', multipart={'file': {'name':'layout.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':data.getvalue()}}).json()
            task = page.request.post(base+'/api/tasks', data={'title':'Inline edit', 'mode':'quick', 'text':imported['text'], 'source_id':imported['source_id']}).json()
            page.request.post(base+'/api/tasks/'+task['id']+'/review/stream', data={})
            record = workspace_api.repo.get(task['id'], 'task')
            run = workspace_api.repo.get(record['active_run'], 'run')
            result = run['result']
            result['issues'][0]['evidence'] = 'Payment: 100 USD'
            result['issues'][0]['redline'] = None
            workspace_api.repo.update(run['id'], {'result': result})
            page.goto(base+'/workbench?task='+task['id'])
            frame = page.frame_locator('iframe[title="合同原件预览"]')
            expect(frame.locator('article p').first).to_contain_text('Payment: 100 USD')
            close_opinions(page)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            paragraph = frame.locator('[contenteditable="plaintext-only"]').first
            expect(paragraph).to_be_visible()
            paragraph.fill('Discard this edit')
            frame.get_by_role('button', name='取消编辑', exact=True).click()
            expect(frame.locator('article p').first).to_contain_text('Payment: 100 USD')
            close_opinions(page)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            paragraph.fill('Payment: 100 USD222')
            frame.get_by_role('button', name='保存正文', exact=True).click()
            expect(frame.locator('article p').first).to_contain_text('Payment: 100 USD222')
            expect(frame.locator('#save-edit')).to_be_hidden()
            page.reload()
            open_opinions(page)
            expect(frame.locator('article p').first).to_contain_text('Payment: 100 USD222')
            expect(frame.locator('article table')).to_contain_text('Delivery terms')
            # Stale review results must not block repeated manual editing.
            open_opinions(page)
            edit_clause = page.get_by_role('button', name='编辑条款', exact=True).first
            expect(edit_clause).to_be_enabled()
            for value in ['Payment: 100 USD222333', 'Payment: 100 USD222333444']:
                edit_clause.click()
                editor = page.locator('.fixed textarea').last
                expect(editor).to_have_value('Payment: 100 USD222' if value == 'Payment: 100 USD222333' else 'Payment: 100 USD222333')
                editor.fill(value)
                page.get_by_role('button', name='保存条款', exact=True).click()
                expect(frame.locator('article p').first).to_contain_text(value)
                expect(frame.locator('.lexis-located')).to_have_text(value)
                page.reload()
                open_opinions(page)
                expect(frame.locator('article p').first).to_contain_text(value)
            expect(edit_clause).to_be_enabled()
            close_opinions(page)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            frame.locator('[contenteditable="plaintext-only"]').last.fill('Delivery terms updated')
            frame.get_by_role('button', name='保存正文', exact=True).click()
            located = frame.locator('.lexis-located')
            expect(located).to_have_text('Delivery terms updated')
            expect(frame.locator('#locate-status')).to_have_text('已定位到刚修改的位置')
            assert located.evaluate('(el) => { const r=el.getBoundingClientRect(); return r.top >= 0 && r.bottom <= innerHeight; }')
            open_opinions(page)
            held = []
            def delay_first_handling(route):
                if not held:
                    held.append(route)
                else:
                    route.continue_()
            page.route('**/api/runs/*/handling', delay_first_handling)
            with page.expect_request('**/api/runs/*/handling'):
                page.get_by_role('button', name='不采纳', exact=True).first.click()
            page.get_by_role('button', name='待复核', exact=True).first.click()
            response = page.request.fetch(held[0].request)
            held[0].fulfill(response=response)
            close_opinions(page)
            page.get_by_role('button', name='保存', exact=True).click()
            open_opinions(page)
            expect(page.get_by_text('工作内容已保存', exact=True)).to_be_visible()
            card = page.locator('#card-'+result['issues'][0]['id'])
            expect(card.get_by_text('处理状态：待复核', exact=True)).to_be_visible()
            saved = page.request.get(base+'/api/tasks/'+task['id']).json()['run']
            assert saved['decisions'][result['issues'][0]['id']] == 'needs_review'
            page.unroute('**/api/runs/*/handling', delay_first_handling)
            for label, state in [('不采纳', 'rejected'), ('待复核', 'needs_review'), ('重置状态', 'pending')]:
                button = page.get_by_role('button', name=label, exact=True).first
                expect(button).to_be_enabled()
                with page.expect_response(lambda response: '/handling' in response.url and response.request.method == 'PUT'):
                    button.click()
                expect(button).to_have_attribute('aria-pressed', 'true')
                saved = page.request.get(base+'/api/tasks/'+task['id']).json()['run']
                assert saved['decisions'][result['issues'][0]['id']] == state
                page.reload()
                open_opinions(page)
                expect(page.get_by_role('button', name=label, exact=True).first).to_be_enabled()


            close_opinions(page)
            expect(page.get_by_role('button', name='导出', exact=True)).to_be_enabled()
            with page.expect_download() as download:
                page.get_by_role('button', name='导出', exact=True).click()
            assert download.value.suggested_filename.endswith('.docx')
            page.keyboard.press('Escape')
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='审查意见书', exact=True).click()
            expect(page.locator('#printable-report')).to_contain_text('后续编辑尚未重新审查')
            page.get_by_role('button', name='关闭意见书', exact=True).click()
            page.locator('summary').filter(has_text='审查设置').click()
            expect(page.get_by_role('button', name='修改审查条件', exact=True)).to_be_visible()
            expect(page.get_by_role('combobox', name='审查立场', exact=True)).to_be_visible()
            expect(page.get_by_role('button', name='合同正文', exact=True)).to_have_count(0)
            expect(page.get_by_role('button', name='已采纳修订稿', exact=True)).to_have_count(0)
            expect(page.get_by_role('link', name='新建审查', exact=True)).to_have_count(0)
            page.keyboard.press('Escape')
            page.screenshot(path='artifacts/workbench-redesign.png')
            assert not errors, errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        sock.close()


def test_review_activates_original_and_all_suggestions_comparison(isolated_workspace, monkeypatch):
    from io import BytesIO
    from docx import Document
    import asyncio
    monkeypatch.setattr(graph, 'get_llm_client', lambda config: None)
    original_events = workspace_api.review_events
    async def events(inputs, config):
        yield 'progress', {'stage':'analysis_summary','detail':'正在核对付款约定','title':'付款约定核对',
                           'clause':'付款条款','analysis':'付款金额正在结合正文核对，结果尚待完整校验。',
                           'proposal':'核对金额及付款条件。','time':'2026-10-15T00:00:00+00:00','attempt':1,'provisional':True}
        for index in range(20):
            yield 'progress', {'stage':'scope','detail':'条款定位 '+str(index)+'：'+('长条款记录。'*80), 'time':'2026-10-15T00:00:00+00:00'}
        await asyncio.sleep(2)
        async for kind, data in original_events(inputs, config):
            if kind == 'complete':
                data['issues'][0]['evidence'] = 'Payment: 100 USD'
                data['issues'][0]['redline'] = {'change_type':'replace', 'original_text':'Payment: 100 USD', 'suggested_text':'Payment: 200 USD'}
                # Independent unrewritten clause, regardless of rule issues consolidated above.
                data['issues'].append({**data['issues'][0], 'id':'delivery_risk', 'segment_id':None})
                data['issues'][1]['evidence']='Section 20: delivery, payment, acceptance and contract obligations.'
                data['issues'][1]['redline']=None
            yield kind, data
    monkeypatch.setattr(workspace_api, 'review_events', events)
    document = Document()
    document.add_paragraph().add_run('Payment: 100 USD').bold = True
    for index in range(65):
        document.add_paragraph(f'Section {index}: delivery, payment, acceptance and contract obligations.')
    document.add_paragraph('1'*160)
    content = BytesIO()
    document.save(content)
    sock = socket.socket()
    sock.bind(('127.0.0.1', 0))
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets':[sock]}, daemon=True)
    thread.start()
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={'width':1920,'height':1080})
            page.add_init_script("sessionStorage.setItem('LEXIS_REVIEW_DISPLAY','compare')")
            errors=[]
            page.on('pageerror', lambda error: errors.append(str(error)))
            base=f'http://127.0.0.1:{sock.getsockname()[1]}'
            imported=page.request.post(base+'/api/contracts/import', multipart={'file':{'name':'compare.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':content.getvalue()}}).json()
            task=page.request.post(base+'/api/tasks', data={'title':'Document comparison','mode':'quick','text':imported['text'],'source_id':imported['source_id']}).json()
            page.goto(base+'/workbench?task='+task['id'])
            left=page.frame_locator('iframe[title="原始合同文档"]')
            right=page.frame_locator('iframe[title="合同原件预览"]')
            expect(page.get_by_test_id('review-summary')).to_be_visible()
            expect(page.locator('#opinion-drawer')).to_be_visible()
            expect(page.get_by_role('button', name='审查意见面板', exact=True)).to_have_count(0)
            assert page.locator('.prototype-navbar').bounding_box()['height'] <= 60
            expect(page.get_by_test_id('original-document-pane')).to_have_count(0)
            expect(right.locator('article p').first).to_have_text('Payment: 100 USD')
            held_original=[]
            def delay_original(route):
                if not held_original: held_original.append(route)
                else: route.continue_()
            for label in ['开始审查', '重新审查']:
                if label=='开始审查':page.route('**/api/sources/*/file',delay_original)
                else:
                    page.get_by_role('button',name='收起审查概览',exact=True).click()
                    page.get_by_role('button',name='收起审查意见与修订清单',exact=True).click()
                page.get_by_role('button', name=label, exact=True).click()
                expect(page.locator('[data-public-analysis]')).to_contain_text('付款金额正在结合正文核对')
                if label=='开始审查':
                    page.set_viewport_size({'width':1366,'height':600})
                    expect(page.locator('.review-execution')).to_have_attribute('open', '')
                    expect(page.locator('.review-stage-list li')).to_have_count(7)
                    bounds=page.locator('.review-dialog').bounding_box()
                    assert bounds['y']>=10 and bounds['y']+bounds['height']<=590
                    assert page.locator('[data-review-trace]').bounding_box()['height']<=161
                    page.locator('.review-dialog-body').evaluate('(el)=>el.scrollTop=el.scrollHeight')
                    expect(page.get_by_role('button',name='取消审查',exact=True)).to_be_in_viewport()
                    page.screenshot(path='artifacts/review-processing.png')
                    page.set_viewport_size({'width':1920,'height':1080})
                page.get_by_role('button',name='查看审查修订稿',exact=True).click()
                expect(page.locator('.review-dialog')).to_have_count(0)
                expect(page.locator('#review-summary-content')).to_be_visible()
                expect(page.locator('#review-opinion-content')).to_be_visible()
                expect(page.get_by_role('button',name='收起审查概览',exact=True).locator('svg')).to_have_count(1)
                if label=='开始审查':
                    expect(right.locator('article p').first).to_have_text('Payment: 200 USD',timeout=30000)
                    original_response=page.request.fetch(held_original[0].request)
                    held_original[0].fulfill(response=original_response)
                    page.unroute('**/api/sources/*/file',delay_original)
                expect(page.get_by_test_id('original-document-pane')).to_be_visible()
                expect(left.locator('article p').first).to_have_text('Payment: 100 USD')
                expect(right.locator('article p').first).to_have_text('Payment: 200 USD', timeout=30000)
                expect(right.locator('#pane-badge')).to_have_text('全部建议')
                expect(right.locator('#pane-badge')).to_have_attribute('title', '全部修改建议，未自动采纳')
                expect(left.locator('#pane-badge')).to_have_text('只读')
                expect(left.locator('article p').first).to_have_class(re.compile('lexis-review-risk'))
                expect(right.locator('article p').first).to_have_class(re.compile('lexis-review-risk'))
                expect(left.locator('article p').nth(21)).to_have_class(re.compile('lexis-review-risk'))
                expect(right.locator('article p').nth(21)).to_have_class(re.compile('lexis-review-risk'))
                expect(page.locator('.wb-pane-heading')).to_have_count(0)
                assert right.locator('header').bounding_box()['height'] <= 48
                assert left.locator('header').bounding_box()['height'] <= 48
                record=page.request.get(base+'/api/tasks/'+task['id']).json()
                assert record['text']==imported['text']
                assert record['run']['applied']==[]
            # Reload must reconstruct the original baseline and all-suggestions
            # draft from the saved run, without launching another review.
            page.reload()
            expect(page.get_by_test_id('original-document-pane')).to_be_visible()
            expect(left.locator('article p').first).to_have_text('Payment: 100 USD')
            expect(right.locator('article p').first).to_have_text('Payment: 200 USD')
            expect(page.locator('#review-summary-content')).to_be_visible()
            expect(page.locator('#review-opinion-content')).to_be_visible()
            for preview in [left,right]:
                expect(preview.locator('#zoom-percent')).to_have_text('80%')
            assert page.request.get(base+'/api/tasks/'+task['id']).json()['run']['id']==record['run']['id']
            expect(left.locator('.lexis-diff-original')).to_have_text('Payment: 100 USD')
            expect(right.locator('.lexis-diff-revised')).to_have_text('Payment: 200 USD')
            expect(right.locator('mark.lexis-diff-token')).to_have_text('2')
            expect(left.locator('mark.lexis-diff-token')).to_have_text('1')
            assert right.locator('article p').nth(21).evaluate('(el)=>getComputedStyle(el).backgroundColor') == 'rgba(0, 0, 0, 0)'
            assert left.locator('html').evaluate('(el)=>getComputedStyle(el).scrollbarWidth') == 'none'
            right.locator('body').evaluate('()=>{window.dispatchEvent(new Event("wheel"));window.scrollTo(0,(document.documentElement.scrollHeight-innerHeight)*.3);}')
            left.locator('body').evaluate('()=>new Promise((resolve,reject)=>{setTimeout(()=>reject(new Error("Scroll synchronization timed out")),5000);const check=()=>scrollY/Math.max(1,document.documentElement.scrollHeight-innerHeight)>.2?resolve():requestAnimationFrame(check);check();})')
            assert left.locator('body').evaluate('()=>scrollY/Math.max(1,document.documentElement.scrollHeight-innerHeight)') > .2
            left.locator('body').evaluate('()=>{window.dispatchEvent(new Event("wheel"));window.scrollTo(0,document.documentElement.scrollHeight);}')
            right.locator('body').evaluate('()=>new Promise((resolve,reject)=>{setTimeout(()=>reject(new Error("Scroll synchronization timed out")),5000);const check=()=>scrollY/Math.max(1,document.documentElement.scrollHeight-innerHeight)>.95?resolve():requestAnimationFrame(check);check();})')
            assert right.locator('body').evaluate('()=>scrollY/Math.max(1,document.documentElement.scrollHeight-innerHeight)') > .95
            right.locator('body').evaluate('()=>{window.dispatchEvent(new Event("wheel"));window.scrollTo(0,0);}')
            expect(page.get_by_test_id('review-summary')).to_be_visible()
            close_opinions(page)
            before = page.get_by_test_id('revised-document-pane').bounding_box()
            open_opinions(page)
            expect(page.locator('#opinion-drawer')).to_be_visible()
            expect(page.get_by_test_id('review-summary')).to_be_visible()
            revised_box=page.get_by_test_id('revised-document-pane').bounding_box()
            panel_box=page.locator('#opinion-drawer').bounding_box()
            assert revised_box['width'] == before['width']
            assert revised_box['x']+revised_box['width'] <= panel_box['x']
            assert panel_box['x']+panel_box['width'] <= 1920
            assert page.get_by_test_id('review-summary').bounding_box()['width'] >= 320
            assert page.locator('#opinion-drawer').bounding_box()['width'] == 380
            for panel_name,button_name in [('review-summary-panel','收起审查概览'),('opinion-drawer','收起审查意见与修订清单')]:
                panel=page.locator('#'+panel_name).bounding_box()
                handle=page.get_by_role('button',name=button_name,exact=True).bounding_box()
                assert abs(handle['y']+handle['height']/2-panel['y']-panel['height']/2)<2
            action_row=page.locator('.wb-issue-actions').first
            positions=[button.bounding_box()['y'] for button in action_row.get_by_role('button').all()]
            assert max(positions)-min(positions)<2
            actions=action_row.get_by_role('button').all()
            edit_box,ask_box=actions[-2].bounding_box(),actions[-1].bounding_box()
            assert 0 <= ask_box['x']-edit_box['x']-edit_box['width'] <= 8
            page.screenshot(path='artifacts/workbench-docked-opinions.png')
            for width in [1366, 1920]:
                page.set_viewport_size({'width':width,'height':1080})
                page.wait_for_function("""() => {
                    const doc=document.querySelector('[data-testid="revised-document-pane"]').getBoundingClientRect();
                    const panel=document.querySelector('#opinion-drawer').getBoundingClientRect();
                    return doc.right <= panel.left && panel.right <= innerWidth;
                }""")
                for preview in [left, right]:
                    expect(preview.locator('section.docx').first).to_be_visible()
                    expect(preview.locator('#zoom-percent')).to_have_text('80%')
            expect(page.get_by_role('button', name='审查概览', exact=True)).to_have_count(0)
            page.get_by_role('button', name='收起审查概览', exact=True).click()
            expect(page.locator('#review-summary-content')).to_be_hidden()
            assert page.get_by_test_id('review-summary').bounding_box()['width'] == 36
            for preview in [left, right]:
                expect(preview.locator('#zoom-percent')).to_have_text('100%')
            page.get_by_role('button', name='收起审查意见与修订清单', exact=True).click()
            expect(page.locator('#review-opinion-content')).to_be_hidden()
            assert page.locator('#opinion-drawer').bounding_box()['width'] == 36
            expect(right.locator('#zoom-percent')).to_have_text('100%')
            page.get_by_role('button', name='展开审查概览', exact=True).click()
            expect(page.locator('#review-summary-content')).to_be_visible()
            expect(page.locator('#review-opinion-content')).to_be_hidden()
            for preview in [left, right]:
                expect(preview.locator('#zoom-percent')).to_have_text('100%')
            page.get_by_role('button', name='展开审查意见与修订清单', exact=True).click()
            expect(page.locator('#review-opinion-content')).to_be_visible()
            expect(page.locator('#review-summary-content')).to_be_visible()
            page.keyboard.press('Escape')
            expect(page.locator('#opinion-drawer')).to_be_visible()
            expect(page.get_by_role('button', name='收起审查意见面板', exact=True)).to_have_count(0)
            for more in page.locator('.revision-more-actions').all():
                assert more.get_attribute('open') is None
            with page.expect_download() as download:
                page.get_by_role('button', name='导出', exact=True).click()
            assert Document(download.value.path()).paragraphs[0].text == 'Payment: 200 USD'
            for button in page.get_by_role('button', name='不采纳', exact=True).all():
                button.click()
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='完成复核', exact=True).click()
            page.locator('.prototype-document-menu > summary').click()
            expect(page.get_by_role('button', name='已完成复核', exact=True)).to_be_visible()
            assert page.request.get(base+'/api/tasks/'+task['id']).json()['status'] == 'completed'
            page.locator('.prototype-document-menu > summary').click()
            expect(right.locator('#zoom-percent')).to_have_text('80%')
            right.get_by_role('button', name='选择文档缩放比例').click()
            right.get_by_role('menuitemradio', name='125%', exact=True).click()
            expect(right.locator('#zoom-percent')).to_have_text('125%')
            expect(right.locator('#zoom-menu')).to_be_hidden()
            right.get_by_role('button', name='选择文档缩放比例').click()
            right.get_by_role('menuitemradio', name='适应宽度', exact=True).click()
            expect(right.locator('#zoom-percent')).to_have_text('适应宽度')
            page.screenshot(path='artifacts/workbench-dual-document.png')
            close_opinions(page)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            right.locator('[contenteditable="plaintext-only"]').first.fill('Payment: 300 USD')
            right.get_by_role('button', name='保存正文', exact=True).click()
            expect(right.locator('article p').first).to_have_text('Payment: 300 USD')
            expect(left.locator('article p').first).to_have_text('Payment: 100 USD')
            expect(right.locator('.lexis-located')).to_have_text('Payment: 300 USD')
            expect(right.locator('.lexis-diff-revised')).to_have_text('Payment: 300 USD')
            expect(left.locator('.lexis-diff-original')).to_have_text('Payment: 100 USD')
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            right.locator('[contenteditable="plaintext-only"]').first.fill('Payment: 350 USD')
            right.locator('[contenteditable="plaintext-only"]').nth(1).fill('Section 0: body edit must survive')
            expect(right.locator('.lexis-diff-revised')).to_have_count(2)
            expect(left.locator('.lexis-diff-original')).to_have_count(2)
            page.get_by_role('button', name='编辑条款', exact=True).first.click()
            clause_editor=page.locator('.fixed textarea').last
            expect(clause_editor).to_have_value('Payment: 350 USD')
            clause_editor.fill('Payment: 400 USD')
            page.get_by_role('button', name='保存条款', exact=True).click()
            expect(right.locator('article p').first).to_have_text('Payment: 400 USD')
            expect(right.locator('article p').nth(1)).to_have_text('Section 0: body edit must survive')
            expect(right.locator('.lexis-diff-revised')).to_have_count(2)
            expect(left.locator('.lexis-diff-original')).to_have_count(2)
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            right.locator('[contenteditable="plaintext-only"]').nth(2).fill('Temporary text to discard')
            expect(right.locator('.lexis-diff-revised')).to_have_count(3)
            right.get_by_role('button', name='取消编辑', exact=True).click()
            expect(right.locator('.lexis-diff-revised')).to_have_count(2)
            expect(right.locator('article p').nth(2)).not_to_have_text('Temporary text to discard')
            page.locator('.prototype-document-menu > summary').click()
            page.get_by_role('button', name='编辑正文', exact=True).click()
            right.locator('[contenteditable="plaintext-only"]').nth(2).fill('Section 1: saved from toolbar')
            page.get_by_role('button', name='保存', exact=True).click()
            expect(right.locator('#save-edit')).to_be_hidden()
            expect(right.locator('article p').nth(2)).to_have_text('Section 1: saved from toolbar')
            page.get_by_role('button', name='编辑条款', exact=True).first.click()
            clause_editor=page.locator('.fixed textarea').last
            expect(clause_editor).to_have_value('Payment: 400 USD')
            clause_editor.fill('Payment: 500 USD')
            page.get_by_role('button', name='保存条款', exact=True).click()
            expect(right.locator('article p').first).to_have_text('Payment: 500 USD')
            expect(right.locator('article p').nth(1)).to_have_text('Section 0: body edit must survive')
            expect(right.locator('article p').nth(2)).to_have_text('Section 1: saved from toolbar')
            expect(right.locator('.lexis-diff-revised')).to_have_count(3)
            page.reload()
            expect(page.get_by_test_id('original-document-pane')).to_be_visible()
            expect(right.locator('article p').first).to_have_text('Payment: 500 USD')
            expect(right.locator('article p').nth(2)).to_have_text('Section 1: saved from toolbar')
            expect(right.locator('article p').nth(1)).to_have_text('Section 0: body edit must survive')
            page.get_by_role('button',name='重新审查',exact=True).click()
            expect(page.locator('.review-dialog')).to_be_visible()
            page.get_by_role('button',name='取消审查',exact=True).click()
            expect(page.locator('.review-dialog')).to_have_count(0)
            expect(right.locator('article p').first).to_have_text('Payment: 500 USD')
            assert not errors, errors
            browser.close()
    finally:
        server.should_exit=True
        thread.join(timeout=10)
        sock.close()
